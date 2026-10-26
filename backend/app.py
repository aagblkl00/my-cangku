from __future__ import annotations

import asyncio
import ast
import base64
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from time import perf_counter
from typing import Any
from collections.abc import Mapping
import unicodedata

from fastapi import FastAPI, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from .automation.playwright_executor import PlaywrightExecutor
from .automation.playwright_recorder import PlaywrightRecorder
from .artifacts import ArtifactStore
from .klein_capabilities import KleinCapabilityStore
from .mission_engine import aggregate_verdict, build_scenarios, build_tapd_draft, classify_reproduction, validate_environment, verify_execution
from .report import write_html_report, write_regression_report
from .store import StudioStore, now_iso
from .runtime_config import ApplicationConfig

CONFIG = ApplicationConfig.load()
CONFIG.ensure_directories()
ROOT = CONFIG.data_dir
store = StudioStore(CONFIG.database_path)
store.recover_incomplete_runs()
store.ensure_demo_data()
artifacts = ArtifactStore(CONFIG.runs_dir)
klein_capabilities = KleinCapabilityStore(CONFIG.data_dir)
app = FastAPI(title=CONFIG.app_name, version=CONFIG.version)
app.mount("/artifacts", StaticFiles(directory=CONFIG.runs_dir, check_dir=False), name="artifacts")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "tauri://localhost",
        "http://tauri.localhost",
        "https://tauri.localhost",
    ],
    allow_origin_regex=r"https?://(?:localhost|127\.0\.0\.1):\d+",
    allow_methods=["*"],
    allow_headers=["*"],
)
clients: set[WebSocket] = set()
active_executors: dict[str, PlaywrightExecutor] = {}
active_recorders: dict[str, PlaywrightRecorder] = {}
active_regressions: dict[str, threading.Event] = {}
active_auto_generations: dict[str, threading.Event] = {}
ai_route_cache: dict[tuple[str, str], tuple[str, str]] = {}
auto_generation_state: dict[str, dict[str, Any]] = {}
active_agents: dict[str, threading.Event] = {}
active_missions: dict[str, threading.Event] = {}
main_loop: asyncio.AbstractEventLoop | None = None


class ProjectIn(BaseModel):
    name: str = Field(min_length=1)
    description: str = ""


class RequirementImportIn(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    module: str = Field(default="", max_length=120)
    content: str = Field(min_length=1, max_length=500_000)
    source_type: str = Field(default="text", pattern="^(text|markdown|word|excel|json|openapi|image|report|defect|link)$")
    source_name: str = Field(default="", max_length=260)
    entities: list[dict[str, Any] | str] = []
    relations: list[dict[str, str]] = []
    rules: list[dict[str, Any] | str] = []
    pending_questions: list[str] = []


class RequirementVersionReviewIn(BaseModel):
    decision: str = Field(pattern="^(activate|reject)$")


class EnvironmentIn(BaseModel):
    project_id: str
    name: str = Field(min_length=1)
    base_url: str = ""
    variables: dict[str, str] = {}
    browser: str = "chromium"
    read_only: bool = True
    allowed_hosts: list[str] = []
    environment_kind: str = Field(default="test", pattern="^(local|test|staging|production)$")
    payment_mode: str = Field(default="disabled", pattern="^(disabled|mock|sandbox)$")


class ElementIn(BaseModel):
    project_id: str
    element_id: str = ""
    page_url: str = ""
    page_title: str = ""
    name: str = "页面元素"
    element_type: str = "element"
    role: str = ""
    locator_type: str = "css"
    locator_value: str = ""
    fallback_locators: list[dict[str, str]] = []
    health_status: str = "未验证"


class FlowIn(BaseModel):
    flow_id: str = ""
    project_id: str
    name: str = "新自动化流程"
    description: str = ""
    engine: str = "playwright"
    version: int = 1
    environment_id: str | None = None
    variables: dict[str, str] = {}
    breakpoints: list[int] = []
    steps: list[dict[str, Any]] = []
    checkpoint_source: str = "manual"


class FlowCheckpointIn(BaseModel):
    source: str = "manual"
    label: str = ""


class RunIn(BaseModel):
    flow_id: str
    environment_id: str | None = None
    browser: str = "chromium"
    from_step: int = 1
    to_step: int | None = None
    breakpoints: list[int] = []
    runtime_variables: dict[str, Any] = {}
    dataset_name: str = ""
    # AI 定位器自愈配置（可选）：{"base_url": "...", "api_key": "...", "model": "..."}
    # 不传时回退到环境变量 LODESTAR_AI_BASE_URL / LODESTAR_AI_KEY / LODESTAR_AI_MODEL；
    # 两边都为空则关闭 AI 自愈，执行行为与改造前完全一致。
    ai: dict[str, Any] = {}


class RecordingIn(BaseModel):
    project_id: str
    url: str
    flow_id: str | None = None


class TestCaseIn(BaseModel):
    case_id: str = ""
    project_id: str
    title: str = Field(min_length=1)
    module: str = ""
    priority: str = "P1"
    case_type: str = "回归测试"
    preconditions: str = ""
    steps: str = ""
    test_data: str = ""
    expected: str = ""
    tags: list[str] = []
    status: str = "未自动化"
    automation_flow_id: str | None = None


class TestCaseBatchDeleteIn(BaseModel):
    project_id: str
    case_ids: list[str] = Field(min_length=1)


class RegressionPlanIn(BaseModel):
    regression_plan_id: str = ""
    project_id: str
    name: str = Field(min_length=1)
    description: str = ""
    environment_id: str | None = None
    browser: str = "chromium"
    stop_on_failure: bool = False
    cases: list[dict[str, Any] | str] = []


class ImportPreviewIn(BaseModel):
    project_id: str
    filename: str
    content_base64: str
    duplicate_strategy: str = "skip"


class ImportConfirmIn(BaseModel):
    project_id: str
    rows: list[dict[str, Any]]
    duplicate_strategy: str = "skip"


class CleanupPreviewIn(BaseModel):
    days: int | None = None
    before: str | None = None


class ProjectPackageIn(BaseModel):
    content_base64: str
    name: str = ""


class AiVerifyIn(BaseModel):
    base_url: str
    api_key: str = ""
    model: str = ""


class AiModelsIn(BaseModel):
    base_url: str
    api_key: str = ""


class AiGenerateFlowIn(BaseModel):
    base_url: str
    api_key: str = ""
    model: str = ""
    environment_id: str | None = None
    start_url: str = ""


class AiBatchIn(BaseModel):
    base_url: str
    api_key: str = ""
    model: str = ""
    only_unbound: bool = True


class AiChatIn(BaseModel):
    base_url: str
    api_key: str = ""
    model: str
    mode: str = Field(default="chat", pattern="^(chat|plan|build|diagnose)$")
    message: str = Field(min_length=1, max_length=12000)
    context: dict[str, Any] = {}
    history: list[dict[str, str]] = []


class KleinMemoryIn(BaseModel):
    key: str = Field(min_length=1, max_length=80)
    value: str = Field(min_length=1, max_length=4000)


class KleinBackgroundIn(BaseModel):
    content_base64: str = Field(min_length=1, max_length=8_000_000)
    mime_type: str = Field(default="image/jpeg", pattern="^image/(jpeg|png|webp)$")


class KleinMemoryStateIn(BaseModel):
    enabled: bool


class KleinSkillInstallIn(BaseModel):
    filename: str = Field(min_length=1, max_length=180)
    content_base64: str = Field(min_length=1)


class KleinSkillStateIn(BaseModel):
    enabled: bool


class KleinDesktopPathIn(BaseModel):
    path: str = ""


class MissionCreateIn(BaseModel):
    project_id: str
    environment_id: str
    objective: str = Field(min_length=2, max_length=4000)
    depth: str = Field(default="standard", pattern="^(quick|standard|deep|custom)$")
    case_ids: list[str] = []
    time_budget_minutes: int = Field(default=30, ge=1, le=480)
    max_scenarios: int | None = Field(default=None, ge=1, le=100)


class MissionStopIn(BaseModel):
    reason: str = Field(default="用户停止", max_length=200)


@app.get("/api/health")
def health() -> dict[str, Any]:
    try:
        import playwright  # noqa: F401
        playwright_installed = True
    except ImportError:
        playwright_installed = False
    browser_ready = False
    if playwright_installed:
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as pw:
                browser_ready = Path(getattr(pw, CONFIG.default_browser).executable_path).exists()
        except Exception:
            browser_ready = False
    return {"ok": True, "service": CONFIG.app_name, "version": CONFIG.version, "production": CONFIG.production,
            "playwright_installed": playwright_installed, "playwright_status": "ready" if playwright_installed else "missing",
            "browser_status": "ready" if browser_ready else "not_initialized", "database_status": "ready" if store.path.exists() else "missing",
             "database": str(store.path), "data_dir": str(CONFIG.data_dir), "port": CONFIG.backend_port}


@app.post("/api/diagnostics/self-test")
def diagnostics_self_test() -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    started = perf_counter()

    def record(key: str, label: str, ok: bool, detail: str, duration_ms: int = 0) -> None:
        checks.append({"key": key, "label": label, "ok": ok, "detail": detail, "duration_ms": duration_ms})

    step_started = perf_counter()
    try:
        with store._connect() as conn:
            integrity = str(conn.execute("PRAGMA quick_check").fetchone()[0])
            conn.execute("SELECT 1").fetchone()
        record("database", "SQLite 数据库", integrity.lower() == "ok", f"完整性检查：{integrity}", round((perf_counter() - step_started) * 1000))
    except Exception as exc:  # noqa: BLE001
        record("database", "SQLite 数据库", False, str(exc), round((perf_counter() - step_started) * 1000))

    step_started = perf_counter()
    probe = CONFIG.data_dir / ".klein-write-probe"
    try:
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
        record("storage", "用户数据目录", True, f"可写：{CONFIG.data_dir}", round((perf_counter() - step_started) * 1000))
    except Exception as exc:  # noqa: BLE001
        probe.unlink(missing_ok=True)
        record("storage", "用户数据目录", False, str(exc), round((perf_counter() - step_started) * 1000))

    step_started = perf_counter()
    client_asset = CONFIG.program_dir / "Lodestar.exe" if CONFIG.production else CONFIG.frontend_dist / "index.html"
    if CONFIG.production and not client_asset.is_file():
        client_asset = CONFIG.program_dir / "Klein.exe"
    record("frontend", "客户端资源", client_asset.is_file(), str(client_asset), round((perf_counter() - step_started) * 1000))

    step_started = perf_counter()
    try:
        capabilities_ready = klein_capabilities.state_path.is_file() and klein_capabilities.skills_dir.is_dir()
        record("capabilities", "Lodestar 能力存储", capabilities_ready, f"记忆、Skill 与授权：{klein_capabilities.root}", round((perf_counter() - step_started) * 1000))
    except Exception as exc:  # noqa: BLE001
        record("capabilities", "Lodestar 能力存储", False, str(exc), round((perf_counter() - step_started) * 1000))

    step_started = perf_counter()
    local_answer = _local_ai_answer("你是谁", "chat", {})
    record("ai", "Lodestar 本地 AI", "Lodestar" in local_answer, local_answer, round((perf_counter() - step_started) * 1000))

    step_started = perf_counter()
    diagnostic_root = CONFIG.data_dir / ".diagnostics"
    run_id = f"self-test-{uuid.uuid4().hex[:8]}"
    try:
        executor = PlaywrightExecutor(diagnostic_root)
        flow = {"flow_id": "klein-self-test", "variables": {}, "elements": {}, "steps": [
            {"id": "self-open", "name": "打开本地健康接口", "action": "open", "value": f"http://127.0.0.1:{CONFIG.backend_port}/api/health", "risk": "read"},
            {"id": "self-assert", "name": "校验本地服务地址", "action": "assert_url", "expected": "/api/health", "risk": "read"},
        ]}
        environment = {"browser": CONFIG.default_browser, "headless": True, "read_only": True, "allowed_hosts": ["127.0.0.1"], "variables": {}}
        status, results, _ = executor.run(flow, environment, run_id)
        ok = status == "PASS" and len(results) == 2 and all(item.get("status") == "PASS" for item in results)
        record("browser", "自动化执行链路", ok, f"Chromium · {len(results)}/2 步通过 · {status}", round((perf_counter() - step_started) * 1000))
    except Exception as exc:  # noqa: BLE001
        record("browser", "自动化执行链路", False, str(exc)[:300], round((perf_counter() - step_started) * 1000))
    finally:
        shutil.rmtree(diagnostic_root / run_id, ignore_errors=True)

    return {"ok": all(item["ok"] for item in checks), "checks": checks, "duration_ms": round((perf_counter() - started) * 1000), "checked_at": now_iso()}


@app.post("/api/ai/ui-actions")
def validate_ai_ui_actions(payload: dict[str, Any]) -> dict[str, Any]:
    from .ai_ui_actions import normalize_ui_action
    try:
        actions = [normalize_ui_action(item) for item in payload.get("actions", [])]
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, str(exc)) from exc
    return {"actions": actions}


@app.post("/api/runtime/browser/init")
def initialize_browser() -> dict[str, Any]:
    """Initialize Chromium into the user's data directory without requiring a command prompt."""
    if CONFIG.default_browser != "chromium":
        raise HTTPException(400, "当前仅支持初始化 Chromium 浏览器组件")
    env = os.environ.copy()
    env["PLAYWRIGHT_BROWSERS_PATH"] = str(CONFIG.playwright_browser_dir)
    try:
        command = [sys.executable, "--install-browser"] if getattr(sys, "frozen", False) else [sys.executable, "-m", "playwright", "install", "chromium"]
        completed = subprocess.run(command, env=env, capture_output=True, text=True, timeout=600, check=False)
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(504, "浏览器组件初始化超时") from exc
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "浏览器安装失败").strip()[-1000:]
        raise HTTPException(500, f"浏览器组件初始化失败：{detail}")
    return {"status": "ready", "browser": "chromium", "path": str(CONFIG.playwright_browser_dir)}


@app.get("/api/projects")
def projects() -> list[dict[str, Any]]:
    return store.list_projects()


@app.post("/api/projects")
def create_project(payload: ProjectIn) -> dict[str, Any]:
    return store.create_project(payload.name, payload.description)


@app.delete("/api/projects/{project_id}")
def remove_project(project_id: str) -> dict[str, Any]:
    if not store.get_project(project_id):
        raise HTTPException(404, "项目不存在")
    if len(store.list_projects()) <= 1:
        raise HTTPException(400, "至少保留一个项目，最后一个项目不能删除")
    run_ids = store.delete_project(project_id)
    for run_id in run_ids:
        shutil.rmtree(CONFIG.runs_dir / run_id, ignore_errors=True)
    return {"deleted": True, "project_id": project_id}


@app.get("/api/projects/{project_id}/brain")
def project_brain(project_id: str) -> dict[str, Any]:
    try:
        return store.get_project_brain(project_id)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc


@app.post("/api/projects/{project_id}/brain/requirements")
def import_project_requirement(project_id: str, payload: RequirementImportIn) -> dict[str, Any]:
    try:
        return store.import_requirement(project_id, payload.model_dump())
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.get("/api/requirements/{requirement_id}/versions")
def requirement_versions(requirement_id: str) -> list[dict[str, Any]]:
    if not store.get_requirement(requirement_id):
        raise HTTPException(404, "需求不存在")
    return store.list_requirement_versions(requirement_id)


@app.put("/api/requirement-versions/{version_id}/review")
def review_requirement_version(version_id: str, payload: RequirementVersionReviewIn) -> dict[str, Any]:
    try:
        value = store.review_requirement_version(version_id, payload.decision)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    if not value:
        raise HTTPException(404, "需求版本不存在")
    return value


def _redact_package(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: ("${PASSWORD}" if any(token in key.lower() for token in ("password", "token", "secret", "apikey", "api_key")) else _redact_package(item)) for key, item in value.items()}
    if isinstance(value, list):
        return [_redact_package(item) for item in value]
    return value


def _decode_export_rows(rows: list[dict[str, Any]], columns: tuple[str, ...]) -> list[dict[str, Any]]:
    decoded: list[dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        for column in columns:
            if column in item and isinstance(item[column], str):
                try:
                    item[column] = json.loads(item[column])
                except json.JSONDecodeError:
                    pass
        decoded.append(item)
    return decoded


@app.get("/api/projects/{project_id}/export")
def export_project(project_id: str) -> Response:
    project = store.get_project(project_id)
    if not project:
        raise HTTPException(404, "项目不存在")
    with store._connect() as conn:  # export is a read-only snapshot of the existing SQLite data
        environments = _decode_export_rows([dict(row) for row in conn.execute("SELECT * FROM environments WHERE project_id=?", (project_id,)).fetchall()], ("variables_json", "allowed_hosts_json"))
        elements = _decode_export_rows([dict(row) for row in conn.execute("SELECT * FROM elements WHERE project_id=?", (project_id,)).fetchall()], ("fallback_locators_json",))
        flows = [store.get_flow(row["flow_id"]) for row in conn.execute("SELECT flow_id FROM flows WHERE project_id=?", (project_id,)).fetchall()]
        cases = store.list_test_cases(project_id)
        plans = store.list_regression_plans(project_id)
    payload = _redact_package({"format": "automationstudio", "version": 1, "project": project, "environments": environments,
                               "elements": elements, "flows": flows, "test_cases": cases, "regression_plans": plans})
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("project.json", json.dumps(payload, ensure_ascii=False, indent=2))
    import urllib.parse
    filename = f"{project['name']}.automationstudio".replace('\\', '_').replace('/', '_')
    disposition = f'attachment; filename="automationstudio-project.automationstudio"; filename*=UTF-8\'\'{urllib.parse.quote(filename)}'
    return Response(buffer.getvalue(), media_type="application/zip", headers={"Content-Disposition": disposition})


@app.post("/api/projects/import")
def import_project(payload: ProjectPackageIn) -> dict[str, Any]:
    try:
        raw = base64.b64decode(payload.content_base64)
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            data = json.loads(archive.read("project.json").decode("utf-8"))
    except (ValueError, KeyError, json.JSONDecodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, "项目包格式无效") from exc
    source = data.get("project") or {}
    project = store.create_project(payload.name.strip() or f"{source.get('name', '导入项目')}（导入）", source.get("description", ""))
    project_id = project["id"]
    env_map: dict[str, str] = {}
    for item in data.get("environments", []):
        variables = item.get("variables_json", item.get("variables", {}))
        allowed_hosts = item.get("allowed_hosts_json", item.get("allowed_hosts", []))
        if isinstance(variables, str):
            variables = json.loads(variables or "{}")
        if isinstance(allowed_hosts, str):
            allowed_hosts = json.loads(allowed_hosts or "[]")
        created = store.create_environment(project_id, item.get("name", "测试环境"), item.get("base_url", ""), variables,
                                            item.get("browser", "chromium"), bool(item.get("read_only", 1)), allowed_hosts)
        env_map[item.get("id", "")] = created["id"]
    element_map: dict[str, str] = {}
    for item in data.get("elements", []):
        created = store.upsert_element({"project_id": project_id, "name": item.get("name", "页面元素"), "page_url": item.get("page_url", ""),
                                        "page_title": item.get("page_title", ""), "element_type": item.get("element_type", "element"), "role": item.get("role", ""),
                                        "locator_type": item.get("locator_type", "css"), "locator_value": item.get("locator_value", ""),
                                        "fallback_locators": item.get("fallback_locators_json", item.get("fallback_locators", []))})
        element_map[item.get("element_id", "")] = created["element_id"]
    flow_map: dict[str, str] = {}
    for item in data.get("flows", []):
        if not item:
            continue
        steps = []
        for step in item.get("steps", []):
            copied = dict(step)
            if copied.get("element_id") in element_map:
                copied["element_id"] = element_map[copied["element_id"]]
            steps.append(copied)
        created = store.save_flow({"project_id": project_id, "name": item.get("name", "导入流程"), "description": item.get("description", ""),
                                   "engine": item.get("engine", "playwright"), "environment_id": env_map.get(item.get("environment_id")),
                                   "variables": item.get("variables", {}), "breakpoints": item.get("breakpoints", []), "steps": steps})
        flow_map[item.get("flow_id", "")] = created["flow_id"]
    case_map: dict[str, str] = {}
    for item in data.get("test_cases", []):
        if not isinstance(item, Mapping):
            continue
        normalized = _normalize_import_row(item)
        # Test-case IDs are global primary keys in the current schema. Generate
        # a new ID for an imported project and retain the old-to-new mapping so
        # importing a package cannot overwrite a case in the source project.
        created = store.save_test_case({"project_id": project_id,
                                        "title": normalized.get("title") or item.get("title", "导入用例"),
                                        "module": normalized.get("module") or item.get("module", ""),
                                        "priority": _normalize_priority(normalized.get("priority") or item.get("priority", "P1")),
                                        "case_type": normalized.get("case_type") or item.get("case_type", "回归测试"),
                                        "preconditions": normalized.get("preconditions") or item.get("preconditions", ""),
                                        "steps": normalized.get("steps") or item.get("steps", ""),
                                        "test_data": normalized.get("test_data") or item.get("test_data", ""),
                                        "expected": normalized.get("expected") or item.get("expected", ""),
                                        "tags": normalized.get("tags", item.get("tags", [])),
                                        "status": normalized.get("status") or item.get("status", "未自动化"),
                                        "automation_flow_id": flow_map.get(item.get("automation_flow_id"))})
        case_map[item.get("case_id", "")] = created["case_id"]
    imported_plans = 0
    for item in data.get("regression_plans", []):
        created = store.save_regression_plan({"project_id": project_id, "name": item.get("name", "导入回归计划"), "description": item.get("description", ""),
                                              "environment_id": env_map.get(item.get("environment_id")), "browser": item.get("browser", "chromium"),
                                              "stop_on_failure": item.get("stop_on_failure", False), "cases": [{"case_id": case_map[c.get("case_id", "")], "enabled": c.get("enabled", 1)} for c in item.get("cases", []) if c.get("case_id") in case_map]})
        imported_plans += 1 if created else 0
    return {"project": project, "flows": len(flow_map), "elements": len(element_map), "test_cases": len(case_map), "regression_plans": imported_plans}


@app.get("/api/projects/{project_id}/environments")
def environments(project_id: str) -> list[dict[str, Any]]:
    return store.list_environments(project_id)


@app.post("/api/environments")
def create_environment(payload: EnvironmentIn) -> dict[str, Any]:
    return store.create_environment(
        payload.project_id, payload.name, payload.base_url, payload.variables, payload.browser,
        payload.read_only, payload.allowed_hosts, payload.environment_kind, payload.payment_mode,
    )


@app.put("/api/environments/{environment_id}")
def update_environment(environment_id: str, payload: EnvironmentIn) -> dict[str, Any]:
    value = store.update_environment(environment_id, payload.model_dump())
    if not value:
        raise HTTPException(404, "测试环境不存在")
    return value


@app.delete("/api/environments/{environment_id}")
def delete_environment(environment_id: str) -> dict[str, bool]:
    try:
        if not store.delete_environment(environment_id):
            raise HTTPException(404, "测试环境不存在")
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    return {"deleted": True}


def _mission_preflight(mission: dict[str, Any]) -> list[dict[str, Any]]:
    environment = next(
        (item for item in store.list_environments(mission["project_id"]) if item["id"] == mission.get("environment_id")),
        None,
    )
    findings: list[dict[str, Any]] = []
    if not environment:
        findings.append({"code": "ENVIRONMENT_MISSING", "severity": "block", "message": "未找到当前项目的测试环境"})
    else:
        findings.extend(validate_environment(environment))
    if not mission.get("scenarios"):
        findings.append({"code": "NO_SCENARIOS", "severity": "block", "message": "没有匹配到可执行测试场景"})
    for scenario in mission.get("scenarios", []):
        case = store.get_test_case(str(scenario.get("case_id") or ""))
        scenario_findings: list[dict[str, str]] = []
        if not case:
            scenario_findings.append({"code": "CASE_MISSING", "severity": "block", "message": "来源用例不存在"})
        else:
            if not str(case.get("steps") or "").strip() or not str(case.get("expected") or "").strip():
                scenario_findings.append({"code": "CASE_INCOMPLETE", "severity": "block", "message": "用例缺少操作步骤或预期结果"})
            if not case.get("automation_flow_id"):
                scenario_findings.append({"code": "EXECUTION_ADAPTER_MISSING", "severity": "block", "message": "当前阶段尚无可用的 UI 执行适配器"})
        store.update_mission_scenario(scenario["scenario_id"], findings=scenario_findings)
        findings.extend({**item, "scenario_id": scenario["scenario_id"], "scenario_title": scenario["title"]} for item in scenario_findings)
    return findings


@app.get("/api/projects/{project_id}/missions")
def missions(project_id: str, limit: int = 50) -> list[dict[str, Any]]:
    return store.list_missions(project_id, limit)


@app.get("/api/projects/{project_id}/mission-defect-drafts")
def mission_defect_drafts(project_id: str) -> list[dict[str, Any]]:
    drafts: list[dict[str, Any]] = []
    for summary in store.list_missions(project_id, 200):
        mission = store.get_mission(summary["mission_id"]) or {}
        for scenario in mission.get("scenarios", []):
            draft = scenario.get("tapd_draft") or {}
            if not draft.get("title"):
                continue
            drafts.append({
                "runId": scenario.get("reproduction_run_id") or scenario.get("run_id") or scenario["scenario_id"],
                "missionId": mission.get("mission_id", ""), "scenarioId": scenario["scenario_id"],
                "title": draft.get("title", ""), "preconditions": draft.get("preconditions", ""),
                "steps": draft.get("steps", ""), "expected": draft.get("expected", ""),
                "actual": draft.get("actual", ""), "status": "draft",
                "classification": scenario.get("failure_classification", ""),
                "confidence": scenario.get("reproduction_confidence", 0),
            })
    return drafts


@app.get("/api/missions/{mission_id}")
def mission_detail(mission_id: str) -> dict[str, Any]:
    value = store.get_mission(mission_id)
    if not value:
        raise HTTPException(404, "自主测试任务不存在")
    value["active"] = mission_id in active_missions
    return value


@app.post("/api/missions")
def create_mission(payload: MissionCreateIn) -> dict[str, Any]:
    if not store.get_project(payload.project_id):
        raise HTTPException(404, "测试项目不存在")
    environment = next((item for item in store.list_environments(payload.project_id) if item["id"] == payload.environment_id), None)
    if not environment:
        raise HTTPException(400, "请选择当前项目的测试环境")
    cases = store.list_test_cases(payload.project_id)
    scenarios = build_scenarios(cases, payload.objective, payload.depth, payload.case_ids, payload.max_scenarios)
    value = store.create_mission({
        "project_id": payload.project_id,
        "environment_id": payload.environment_id,
        "objective": payload.objective.strip(),
        "depth": payload.depth,
        "case_ids": payload.case_ids,
        "config": {"time_budget_minutes": payload.time_budget_minutes, "max_scenarios": payload.max_scenarios},
    }, scenarios)
    store.update_mission(value["mission_id"], status="PREFLIGHT", progress=15)
    findings = _mission_preflight(store.get_mission(value["mission_id"]) or value)
    blocked = [item for item in findings if item.get("severity") == "block"]
    store.update_mission(value["mission_id"], blocking_reason="；".join(item["message"] for item in blocked[:5]))
    store.add_mission_event(value["mission_id"], "preflight_finished", "PREFLIGHT",
                            "前置检查完成" if not blocked else f"前置检查发现 {len(blocked)} 个阻塞项",
                            {"findings": findings})
    return store.get_mission(value["mission_id"]) or value


@app.post("/api/missions/{mission_id}/start")
def start_mission(mission_id: str) -> dict[str, Any]:
    value = store.get_mission(mission_id)
    if not value:
        raise HTTPException(404, "自主测试任务不存在")
    if mission_id in active_missions:
        return value
    if value.get("status") != "PREFLIGHT":
        raise HTTPException(409, "当前任务不处于可启动的前置检查阶段")
    findings = _mission_preflight(value)
    blocked = [item for item in findings if item.get("severity") == "block"]
    if blocked:
        reason = "；".join(item["message"] for item in blocked[:5])
        store.update_mission(mission_id, blocking_reason=reason)
        raise HTTPException(409, reason)
    stop_event = threading.Event()
    active_missions[mission_id] = stop_event
    threading.Thread(target=_run_mission, args=(mission_id, stop_event), daemon=True).start()
    return store.get_mission(mission_id) or value


@app.post("/api/missions/{mission_id}/stop")
def stop_mission(mission_id: str, payload: MissionStopIn) -> dict[str, Any]:
    value = store.get_mission(mission_id)
    if not value:
        raise HTTPException(404, "自主测试任务不存在")
    stop_event = active_missions.get(mission_id)
    if not stop_event:
        raise HTTPException(409, "当前任务没有在运行")
    stop_event.set()
    store.add_mission_event(mission_id, "stop_requested", value.get("status", ""), payload.reason)
    return {"mission_id": mission_id, "status": "STOPPING"}


@app.get("/api/projects/{project_id}/elements")
def elements(project_id: str) -> list[dict[str, Any]]:
    return store.list_elements(project_id)


@app.post("/api/elements")
def save_element(payload: ElementIn) -> dict[str, Any]:
    return store.upsert_element(payload.model_dump())


@app.delete("/api/elements/{element_id}")
def delete_element(element_id: str) -> dict[str, bool]:
    if store.get_element_usage(element_id):
        raise HTTPException(409, "页面元素仍被流程引用，不能删除；请先替换引用")
    if not store.delete_element(element_id):
        raise HTTPException(404, "页面元素不存在")
    return {"deleted": True}


@app.get("/api/elements/{element_id}/references")
def element_references(element_id: str) -> list[dict[str, Any]]:
    if not store.get_element(element_id):
        raise HTTPException(404, "页面元素不存在")
    return store.get_element_usage(element_id)


@app.post("/api/elements/{element_id}/verify")
def verify_element(element_id: str) -> dict[str, Any]:
    element = store.get_element(element_id)
    if not element:
        raise HTTPException(404, "页面元素不存在")
    if not element.get("page_url"):
        raise HTTPException(400, "页面元素缺少页面地址，无法验证")
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.goto(element["page_url"], wait_until="domcontentloaded", timeout=15000)
                count = PlaywrightExecutor._locator(page, element).count()
            finally:
                browser.close()
        health = "优秀" if count == 1 else ("警告" if count > 1 else "脆弱")
        return store.mark_element_verified(element_id, health) or element
    except Exception as exc:  # noqa: BLE001 - verification is an explicit diagnostic action.
        store.mark_element_verified(element_id, "脆弱")
        raise HTTPException(422, f"元素验证失败：{exc}") from exc


def _safe_calculate(expression: str) -> int | float | None:
    """Evaluate a small arithmetic expression without exposing Python eval."""
    operators = {
        ast.Add: lambda left, right: left + right,
        ast.Sub: lambda left, right: left - right,
        ast.Mult: lambda left, right: left * right,
        ast.Div: lambda left, right: left / right,
        ast.FloorDiv: lambda left, right: left // right,
        ast.Mod: lambda left, right: left % right,
        ast.Pow: lambda left, right: left ** right,
        ast.UAdd: lambda value: value,
        ast.USub: lambda value: -value,
    }

    def calculate(node: ast.AST) -> int | float:
        if isinstance(node, ast.Expression):
            return calculate(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in operators:
            left = calculate(node.left)
            right = calculate(node.right)
            if isinstance(node.op, ast.Pow) and abs(right) > 10:
                raise ValueError("exponent is too large")
            return operators[type(node.op)](left, right)
        if isinstance(node, ast.UnaryOp) and type(node.op) in operators:
            return operators[type(node.op)](calculate(node.operand))
        raise ValueError("unsupported expression")

    candidate = expression.strip().replace("×", "*").replace("÷", "/")
    if not candidate or len(candidate) > 80 or not re.fullmatch(r"[\d\s+\-*/%.()]+", candidate):
        return None
    try:
        value = calculate(ast.parse(candidate, mode="eval"))
        if not isinstance(value, (int, float)) or abs(value) > 10**15:
            return None
        return round(value, 8) if isinstance(value, float) else value
    except (ArithmeticError, SyntaxError, TypeError, ValueError):
        return None


def _analyze_desktop_document(name: str, content: str) -> str:
    cleaned = re.sub(r"<!--.*?-->", "", content, flags=re.S)
    cleaned = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", cleaned)
    cleaned = re.sub(r"<br\s*/?>", "\n", cleaned, flags=re.I)
    cleaned = re.sub(r"<[^>]+>", "", cleaned)
    lines = [re.sub(r"\s+", " ", line).strip(" \t-+*") for line in cleaned.splitlines()]
    lines = [line for line in lines if len(line) >= 2]
    headings: list[str] = []
    for match in re.finditer(r"^#{1,4}\s+(.+)$", content, flags=re.M):
        heading = re.sub(r"<[^>]+>", "", match.group(1)).strip()
        if heading and heading not in headings:
            headings.append(heading)
    signal_words = ("必须", "需要", "新增", "修改", "特殊", "不允许", "不计算", "排除", "支持", "规则", "异常", "默认", "统一", "逻辑", "注意")
    rules: list[str] = []
    for line in lines:
        if any(word in line for word in signal_words) and line not in rules:
            rules.append(line[:180])
        if len(rules) >= 10:
            break
    dimensions = []
    keyword_dimensions = [
        (("订单", "支付", "退款"), "订单状态、金额口径、退款与重复处理"),
        (("用户", "学员", "手机号"), "用户合并、手机号脱敏、身份与历史数据"),
        (("标签", "商品", "分组"), "商品标签映射、缺失标签、分组与筛选"),
        (("权限", "审核", "角色"), "角色权限、审核流与越权访问"),
        (("时间", "日期", "到期"), "时间边界、历史规则与权益有效期"),
        (("导入", "导出", "批量"), "批量数据、导入校验、导出一致性与幂等"),
    ]
    for keywords, description in keyword_dimensions:
        if any(keyword in cleaned for keyword in keywords):
            dimensions.append(description)
    if not dimensions:
        dimensions = ["主流程、边界值、异常恢复与数据一致性", "权限、重复提交和状态回退", "失败证据与回归范围"]
    sections = "\n".join(f"- {item}" for item in headings[:8]) or "- 未识别到明确章节，建议先补充分段标题"
    key_rules = "\n".join(f"{index}. {item}" for index, item in enumerate(rules, 1)) or "1. 暂未提取到明确规则句，建议人工确认需求目标与验收口径。"
    test_focus = "\n".join(f"- {item}" for item in dimensions[:6])
    return (
        f"已读取并分析《{name}》。\n\n"
        f"文档结构\n{sections}\n\n"
        f"关键规则与风险线索\n{key_rules}\n\n"
        f"建议优先验证\n{test_focus}\n\n"
        "以上是基于文档原文的快速提取。涉及金额口径、历史数据迁移和状态规则的结论，仍应结合实际页面、接口或数据库证据复核。"
    )


def _local_skill_answer(message: str, skills: Any) -> str:
    """Apply simple text-response rules from installed Skills in offline mode."""
    if not isinstance(skills, list):
        return ""
    normalized_message = re.sub(r"\s+", "", message).casefold()
    if not normalized_message:
        return ""
    patterns = (
        re.compile(
            r"当用户(?:询问|问到|输入|说出)?\s*[“\"']?(.{1,80}?)[”\"']?\s*时[，,\s]*"
            r"(?:回答|回复|输出)\s*[：:]\s*[`“\"']?(.{1,500}?)[`”\"']?(?:[。\n]|$)",
            re.I,
        ),
        re.compile(
            r"when\s+(?:the\s+)?user\s+(?:asks|says|types)\s*[\"']?(.{1,80}?)[\"']?\s*[,，]?\s*"
            r"(?:reply|answer|respond)(?:\s+with)?\s*[：:]?\s*[`\"']?(.{1,500}?)[`\"']?(?:[.\n]|$)",
            re.I,
        ),
    )
    for skill in skills[:20]:
        if not isinstance(skill, Mapping):
            continue
        instructions = str(skill.get("instructions") or "")[:4000]
        for pattern in patterns:
            for match in pattern.finditer(instructions):
                trigger = re.sub(r"\s+", "", match.group(1)).casefold()
                if trigger and trigger in normalized_message:
                    return match.group(2).strip().strip("`'\"“” ")
    return ""


def _local_ai_answer(message: str, mode: str, context: dict[str, Any]) -> str:
    """Provide useful offline QA assistance when no external model is configured."""
    text = message.strip()
    lowered = text.lower()
    project = str(context.get("project") or context.get("projectName") or "当前工作区")
    view = str(context.get("view") or context.get("activeView") or "当前页面")
    selected = context.get("selectedCase") or context.get("case") or context.get("selectedFlow") or {}
    selected_name = ""
    if isinstance(selected, Mapping):
        selected_name = str(selected.get("title") or selected.get("name") or selected.get("id") or "")

    memories = context.get("Lodestar长期记忆") or context.get("Klein长期记忆") or []
    if any(keyword in text for keyword in ("你记得什么", "长期记忆", "我的记忆")):
        if not memories:
            return "目前还没有长期记忆。你可以直接说“记住：……”或在设置的 Lodestar 能力中心添加。"
        lines = [f"- {item.get('key', '记忆')}：{item.get('value', '')}" for item in memories if isinstance(item, Mapping)]
        return "我记得这些内容：\n\n" + "\n".join(lines[:50])

    skills = context.get("Lodestar启用Skills") or context.get("Klein启用Skills") or []
    skill_answer = _local_skill_answer(text, skills)
    if skill_answer:
        return skill_answer
    if "skill" in lowered or "技能" in text:
        names = [str(item.get("name", "")) for item in skills if isinstance(item, Mapping) and item.get("name")]
        if not names:
            return "目前没有启用的 Skill。你可以在设置的 Lodestar 能力中心安装 Markdown、JSON 或 ZIP Skill 包。"
        return "当前已启用的 Skill：\n\n" + "\n".join(f"- {name}" for name in names)

    desktop_context = context.get("Lodestar桌面上下文") or context.get("Klein桌面上下文")
    if isinstance(desktop_context, Mapping):
        if not desktop_context.get("authorized"):
            return "我还没有 Windows 桌面读取权限。请在设置的 Lodestar 能力中心完成授权。"
        desktop_file = desktop_context.get("file")
        if isinstance(desktop_file, Mapping):
            file_name = str(desktop_file.get("name", "桌面文件"))
            file_content = str(desktop_file.get("content", ""))
            if any(keyword in text for keyword in ("分析", "总结", "提取", "风险", "测试点")):
                return _analyze_desktop_document(file_name, file_content)
            return f"已读取 {file_name}：\n\n{file_content[:12000]}"
        entries = desktop_context.get("entries") or []
        if isinstance(entries, list):
            names = [str(item.get("name", "")) for item in entries if isinstance(item, Mapping) and item.get("name")]
            return "当前桌面可读取项目：\n\n" + "\n".join(f"- {name}" for name in names[:80])

    if mode == "plan":
        objective = text or "完成当前测试闭环"
        plan = {
            "title": "Lodestar 本地测试闭环计划",
            "objective": objective,
            "items": [
                {"title": "确认范围与风险", "detail": "明确核心流程、边界条件、数据和环境前提", "risk": "high"},
                {"title": "组织并执行用例", "detail": "将关键用例加入套件和会话，记录通过、失败、阻塞与跳过", "risk": "high"},
                {"title": "收集执行证据", "detail": "保存步骤、截图、日志、网络和 Trace，避免仅凭结论判断", "risk": "medium"},
                {"title": "失败归因与回归", "detail": "区分产品缺陷、环境问题和脚本问题，失败项进入回归计划", "risk": "high"},
                {"title": "沉淀质量结论", "detail": "复核风险、覆盖率和遗留问题，更新质量视图", "risk": "medium"},
            ],
        }
        return (
            "我已把目标整理为可执行测试闭环；"
            "高风险结论仍需结合真实执行证据人工复核。\n\n"
            "AUTOMATION_STUDIO_PLAN: " + json.dumps(plan, ensure_ascii=False)
        )

    if mode == "diagnose":
        evidence = context.get("recentRun") or context.get("failure") or context.get("evidence")
        evidence_text = json.dumps(evidence, ensure_ascii=False)[:900] if evidence else "当前上下文没有可核验的失败证据"
        return (
            ""
            f"已确认：工作区为“{project}”，页面为“{view}”"
            + (f"，当前对象为“{selected_name}”" if selected_name else "")
            + f"。\n证据状态：{evidence_text}。\n\n"
            "建议按顺序验证：\n"
            "1. 对照失败步骤、实际结果与预期结果，先确认失败是否可稳定复现。\n"
            "2. 检查截图、控制台日志、网络响应和 Trace 的首个异常点。\n"
            "3. 核对测试数据、环境变量、权限和页面前置状态。\n"
            "4. 再区分产品缺陷、环境问题、数据问题或自动化脚本问题。\n\n"
            "在没有证据前，我不会把可能原因直接定性为根因。"
        )

    if mode == "build":
        return (
            ""
            f"针对“{text}”，建议把流程拆为：准备环境与数据、进入目标页面、执行关键动作、"
            "校验业务结果、保存失败证据、清理测试数据。定位器优先使用稳定的 role、label 或 test id，"
            "关键业务结果同时增加页面断言与接口/数据侧复核。当前仅生成设计建议，不会在未确认时直接修改流程。"
        )

    calculation = _safe_calculate(text.rstrip("=?？。！! "))
    if calculation is None:
        arithmetic_candidates = re.findall(
            r"(?<![A-Za-z0-9_])(?:\d+(?:\.\d+)?\s*[+\-*/×÷%]\s*)+\d+(?:\.\d+)?(?![A-Za-z0-9_])",
            text,
        )
        for candidate in arithmetic_candidates:
            calculation = _safe_calculate(candidate)
            if calculation is not None:
                break
    if calculation is not None:
        if any(keyword in text for keyword in ("只回答数字", "只要数字", "仅回答数字")):
            return str(calculation)
        return f"计算结果：{calculation}。这个结果由 Lodestar 本地计算器直接核算，可以确定。"

    if any(keyword in lowered for keyword in ("确定", "确认", "这个数字", "数字对不对", "准确吗")):
        context_numbers = re.findall(r"(?<![A-Za-z])\d+(?:\.\d+)?%?", json.dumps(context, ensure_ascii=False))
        if context_numbers:
            values = "、".join(dict.fromkeys(context_numbers[:8]))
            return (
                f"当前上下文中可读取到的数字是：{values}。"
                "我能确认这些值确实出现在当前页面上下文中，但仅凭上下文不能确认它们的业务口径是否正确；"
                "需要对照数据来源、筛选条件和统计时间范围后才能下最终结论。"
            )
        return "我现在不能确定。当前页面上下文没有提供可核验的数字或数据来源，请把具体数字、统计口径或对应记录发给我后再确认。"

    if any(keyword in lowered for keyword in ("登录", "用例", "测试", "边界", "场景", "覆盖", "验收")):
        return (
            "建议从以下维度检查：\n\n"
            "1. 主流程：有效数据下完整走通，并核对页面、接口和最终业务状态。\n"
            "2. 边界：空值、最小值、最大值、超长值、重复值和临界时间。\n"
            "3. 异常：非法格式、无权限、状态冲突、网络超时、重复提交和服务端失败。\n"
            "4. 状态：刷新、返回、重新进入、并发操作和中断恢复后的数据一致性。\n"
            "5. 证据：每个失败项保留测试数据、操作步骤、预期/实际结果、截图或日志。\n\n"
            f"当前上下文：{project} / {view}"
            + (f" / {selected_name}" if selected_name else "")
            + "。如需理解任意开放问题、文档或图片，请在设置中接入外部模型。"
        )

    if any(keyword in lowered for keyword in ("你好", "hello", "hi")):
        return "你好。"

    identity_prompts = ("你是谁", "你叫什么", "你的名字", "who are you", "what is your name")
    if any(prompt in lowered for prompt in identity_prompts):
        return "Lodestar。"

    return (
        "我已收到这个问题，但当前信息不足以给出可核验的结论。"
        "请补充目标、相关文本或文件，我会继续拆解、分析并给出可执行结果；"
        "需要开放知识问答或图片理解时，也可以在设置中连接你常用的模型服务。"
    )


@app.get("/api/klein/memories")
def klein_memories() -> list[dict[str, Any]]:
    return klein_capabilities.list_memories()


def _klein_background_paths() -> tuple[Path, Path]:
    appearance_dir = klein_capabilities.root / "appearance"
    appearance_dir.mkdir(parents=True, exist_ok=True)
    return appearance_dir, appearance_dir / "background"


@app.get("/api/klein/appearance/background")
def klein_background_state() -> dict[str, Any]:
    _, base = _klein_background_paths()
    candidates = [(base.with_suffix(ext), mime) for ext, mime in ((".jpg", "image/jpeg"), (".png", "image/png"), (".webp", "image/webp"))]
    for path, mime in candidates:
        if path.is_file():
            stat = path.stat()
            return {"enabled": True, "mime_type": mime, "version": str(stat.st_mtime_ns), "path": str(path)}
    return {"enabled": False, "mime_type": "", "version": ""}


@app.get("/api/klein/appearance/background/image")
def klein_background_image() -> Response:
    _, base = _klein_background_paths()
    candidates = [(base.with_suffix(ext), mime) for ext, mime in ((".jpg", "image/jpeg"), (".png", "image/png"), (".webp", "image/webp"))]
    for path, mime in candidates:
        if path.is_file():
            return Response(path.read_bytes(), media_type=mime, headers={"Cache-Control": "no-store"})
    raise HTTPException(404, "未设置 AI 对话背景")


@app.post("/api/klein/appearance/background")
def save_klein_background(payload: KleinBackgroundIn) -> dict[str, Any]:
    try:
        raw = base64.b64decode(payload.content_base64, validate=True)
    except (ValueError, base64.binascii.Error) as exc:
        raise HTTPException(400, "背景图片编码无效") from exc
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(413, "背景图片不能超过 5 MB")
    signatures = {
        "image/jpeg": raw.startswith(b"\xff\xd8\xff"),
        "image/png": raw.startswith(b"\x89PNG\r\n\x1a\n"),
        "image/webp": raw.startswith(b"RIFF") and raw[8:12] == b"WEBP",
    }
    if not signatures.get(payload.mime_type, False):
        raise HTTPException(400, "背景图片格式与内容不匹配")
    _, base = _klein_background_paths()
    suffix = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}[payload.mime_type]
    target = base.with_suffix(suffix)
    temporary = target.with_suffix(target.suffix + ".tmp")
    try:
        temporary.write_bytes(raw)
        temporary.replace(target)
        for old_suffix in (".jpg", ".png", ".webp"):
            old = base.with_suffix(old_suffix)
            if old != target and old.exists():
                old.unlink()
        stat = target.stat()
    except OSError as exc:
        raise HTTPException(500, f"背景图片保存失败：{exc}") from exc
    return {"enabled": True, "mime_type": payload.mime_type, "version": str(stat.st_mtime_ns), "path": str(target)}


@app.delete("/api/klein/appearance/background")
def delete_klein_background() -> dict[str, bool]:
    _, base = _klein_background_paths()
    deleted = False
    for suffix in (".jpg", ".png", ".webp"):
        path = base.with_suffix(suffix)
        if path.exists():
            path.unlink()
            deleted = True
    return {"deleted": deleted}


@app.post("/api/klein/memories")
def save_klein_memory(payload: KleinMemoryIn) -> dict[str, Any]:
    try:
        return klein_capabilities.save_memory(payload.key, payload.value)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.delete("/api/klein/memories/{memory_id}")
def delete_klein_memory(memory_id: str) -> dict[str, bool]:
    if not klein_capabilities.delete_memory(memory_id):
        raise HTTPException(404, "记忆不存在")
    return {"deleted": True}


@app.put("/api/klein/memories/{memory_id}")
def update_klein_memory(memory_id: str, payload: KleinMemoryStateIn) -> dict[str, Any]:
    item = klein_capabilities.set_memory_enabled(memory_id, payload.enabled)
    if not item:
        raise HTTPException(404, "记忆不存在")
    return item


@app.get("/api/klein/skills")
def klein_skills() -> list[dict[str, Any]]:
    return klein_capabilities.list_skills()


@app.post("/api/klein/skills/install")
def install_klein_skill(payload: KleinSkillInstallIn) -> dict[str, Any]:
    try:
        return klein_capabilities.install_skill(payload.filename, payload.content_base64)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.put("/api/klein/skills/{skill_id}")
def update_klein_skill(skill_id: str, payload: KleinSkillStateIn) -> dict[str, Any]:
    item = klein_capabilities.set_skill_enabled(skill_id, payload.enabled)
    if not item:
        raise HTTPException(404, "Skill 不存在")
    return item


@app.delete("/api/klein/skills/{skill_id}")
def delete_klein_skill(skill_id: str) -> dict[str, bool]:
    if not klein_capabilities.delete_skill(skill_id):
        raise HTTPException(404, "Skill 不存在")
    return {"deleted": True}


@app.get("/api/klein/desktop/grants")
def klein_desktop_grants() -> list[dict[str, Any]]:
    return klein_capabilities.list_grants()


@app.post("/api/klein/desktop/grants/default")
def grant_klein_desktop() -> dict[str, Any]:
    try:
        return klein_capabilities.grant_desktop()
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.delete("/api/klein/desktop/grants/{grant_id}")
def revoke_klein_desktop(grant_id: str) -> dict[str, bool]:
    if not klein_capabilities.delete_grant(grant_id):
        raise HTTPException(404, "桌面授权不存在")
    return {"deleted": True}


@app.post("/api/klein/desktop/list")
def list_klein_desktop(payload: KleinDesktopPathIn) -> dict[str, Any]:
    try:
        return klein_capabilities.list_desktop(payload.path)
    except PermissionError as exc:
        raise HTTPException(403, str(exc)) from exc
    except (OSError, ValueError) as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/klein/desktop/read")
def read_klein_desktop(payload: KleinDesktopPathIn) -> dict[str, Any]:
    try:
        return klein_capabilities.read_desktop_text(payload.path)
    except PermissionError as exc:
        raise HTTPException(403, str(exc)) from exc
    except (OSError, ValueError) as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/ai/verify")
def ai_verify(payload: AiVerifyIn) -> dict[str, Any]:
    if payload.base_url.strip().lower() in {"local://lodestar", "local://klein", "local", "builtin://klein"}:
        return {"ok": True, "base_url": "local://lodestar", "model": "lodestar-local", "api_kind": "local", "sample": "本地 AI 已就绪"}
    from .ai_client import verify_connection
    try:
        result = verify_connection(payload.base_url, payload.api_key, payload.model)
        route_key = (payload.base_url.strip().rstrip("/"), payload.model.strip())
        ai_route_cache[route_key] = (str(result.get("base_url") or route_key[0]), str(result["api_kind"]))
        return result
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"AI连接验证失败：{exc}") from exc


@app.post("/api/ai/models")
def ai_models(payload: AiModelsIn) -> dict[str, Any]:
    if payload.base_url.strip().lower() in {"local://lodestar", "local://klein", "local", "builtin://klein"}:
        return {"base_url": "local://lodestar", "models": ["lodestar-local"], "selected": "lodestar-local"}
    from .ai_client import list_models
    try:
        return list_models(payload.base_url, payload.api_key)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"模型列表获取失败：{exc}") from exc


@app.post("/api/ai/chat")
def ai_chat(payload: AiChatIn) -> dict[str, Any]:
    from .ai_client import build_chat_prompt
    from .ai_provider import detect_ai_provider, request_ai_text
    started = perf_counter()

    def runtime(
        *,
        source: str,
        provider: str,
        model: str,
        api_kind: str,
        base_url: str,
        fallback: bool = False,
        error: str = "",
    ) -> dict[str, Any]:
        return {
            "source": source,
            "provider": provider,
            "model": model,
            "api_kind": api_kind,
            "base_url": base_url,
            "duration_ms": max(0, round((perf_counter() - started) * 1000)),
            "fallback": fallback,
            "error": error,
        }

    remembered = klein_capabilities.remember_from_message(payload.message)
    enriched_context = {
        **payload.context,
        "Lodestar长期记忆": klein_capabilities.memory_context(),
        "Lodestar启用Skills": klein_capabilities.skill_context(),
        "本次调用配置": {
            "model": payload.model.strip() or "未配置",
            "mode": payload.mode,
        },
    }
    desktop_context = klein_capabilities.desktop_context(payload.message)
    if desktop_context is not None:
        enriched_context["Lodestar桌面上下文"] = desktop_context
    if remembered and payload.base_url.strip().lower() in {"local://lodestar", "local://klein", "local", "builtin://klein"}:
        local_runtime = runtime(source="local", provider="Lodestar Local", model="lodestar-local", api_kind="local", base_url="local://lodestar")
        return {"answer": f"已记住：{remembered['value']}", "runtime": local_runtime, **local_runtime}
    if payload.base_url.strip().lower() in {"local://lodestar", "local://klein", "local", "builtin://klein"}:
        local_runtime = runtime(source="local", provider="Lodestar Local", model="lodestar-local", api_kind="local", base_url="local://lodestar")
        return {"answer": _local_ai_answer(payload.message, payload.mode, enriched_context), "runtime": local_runtime, **local_runtime}
    history = "\n".join(f"{item.get('role', 'user')}：{item.get('text', '')}" for item in payload.history[-8:])
    prompt = build_chat_prompt(payload.message, {**enriched_context, "最近对话": history})
    route_key = (payload.base_url.strip().rstrip("/"), payload.model.strip())
    preferred_route = ai_route_cache.get(route_key)
    budgets = {
        "chat": (22, 1200, 26),
        "plan": (34, 2400, 40),
        "build": (42, 3200, 48),
        "diagnose": (30, 2200, 36),
    }
    timeout, max_tokens, total_timeout = budgets[payload.mode]
    try:
        response = request_ai_text(
            payload.base_url,
            payload.api_key,
            payload.model,
            prompt,
            timeout=timeout,
            preferred_base_url=preferred_route[0] if preferred_route else None,
            preferred_api_kind=preferred_route[1] if preferred_route else None,
            fallback_timeout=3,
            max_tokens=max_tokens,
            max_retries=0,
            retry_backoff_seconds=0,
            total_timeout=total_timeout,
        )
        ai_route_cache[route_key] = (response.base_url, response.api_kind)
        external_runtime = runtime(
            source="external",
            provider=detect_ai_provider(response.base_url).label,
            model=response.model,
            api_kind=response.api_kind,
            base_url=response.base_url,
        )
        return {"answer": response.text, "runtime": external_runtime, **external_runtime}
    except Exception as exc:  # noqa: BLE001 - keep chat usable when the configured provider is unavailable.
        failure_reason = str(exc)[:500]
        local_runtime = runtime(
            source="local",
            provider="Lodestar Local",
            model="lodestar-local",
            api_kind="local-fallback",
            base_url="local://lodestar",
            fallback=True,
            error=failure_reason,
        )
        return {
            "answer": _local_ai_answer(payload.message, payload.mode, enriched_context),
            "runtime": local_runtime,
            **local_runtime,
        }


@app.post("/api/test-cases/{case_id}/generate-flow")
def generate_flow_from_case(case_id: str, payload: AiGenerateFlowIn) -> dict[str, Any]:
    from .ai_client import build_flow_payload, generate_flow_steps
    case = store.get_test_case(case_id)
    if not case:
        raise HTTPException(404, "测试用例不存在")
    environments = store.list_environments(case["project_id"])
    environment = next((item for item in environments if item["id"] == payload.environment_id), environments[0] if environments else {})
    start_url = payload.start_url or environment.get("base_url", "")
    try:
        steps = generate_flow_steps(payload.base_url, payload.api_key, payload.model, case, start_url)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"AI生成失败：{exc}") from exc
    flow_payload = build_flow_payload(case, steps, start_url, environment.get("id"))
    for element in flow_payload.pop("elements", []):
        store.upsert_element(element)
    flow = store.save_flow(flow_payload)
    return {"flow": flow, "steps": steps}


@app.post("/api/projects/{project_id}/auto-generate-batch")
def auto_generate_batch(project_id: str, payload: AiBatchIn) -> dict[str, Any]:
    if project_id in active_auto_generations:
        raise HTTPException(409, "该项目的批量生成正在进行中")
    if not store.list_test_cases(project_id):
        raise HTTPException(400, "项目没有测试用例")
    active_auto_generations[project_id] = threading.Event()
    threading.Thread(target=_auto_generate_batch, args=(project_id, payload), daemon=True).start()
    return {"started": True, "project_id": project_id}


@app.get("/api/projects/{project_id}/auto-generate-status")
def auto_generate_status(project_id: str) -> dict[str, Any]:
    state = auto_generation_state.get(project_id)
    if not state:
        return {"done": 0, "total": 0, "failed": 0, "current": "", "finished": True, "failures": []}
    return state


@app.post("/api/test-cases/{case_id}/agent-run")
def agent_run(case_id: str, payload: AiGenerateFlowIn) -> dict[str, Any]:
    if not store.get_test_case(case_id):
        raise HTTPException(404, "测试用例不存在")
    if case_id in active_agents:
        raise HTTPException(409, "该用例的智能执行正在进行中")
    active_agents[case_id] = threading.Event()
    threading.Thread(target=_run_agent, args=(case_id, payload), daemon=True).start()
    return {"started": True, "case_id": case_id}


@app.get("/api/projects/{project_id}/flows")
def flows(project_id: str) -> list[dict[str, Any]]:
    return store.list_flows(project_id)


@app.get("/api/flows/{flow_id}")
def flow(flow_id: str) -> dict[str, Any]:
    value = store.get_flow(flow_id)
    if not value:
        raise HTTPException(404, "自动化流程不存在")
    return value


@app.post("/api/flows")
def save_flow(payload: FlowIn) -> dict[str, Any]:
    return store.save_flow(payload.model_dump())


@app.get("/api/flows/{flow_id}/versions")
def flow_versions(flow_id: str) -> list[dict[str, Any]]:
    if not store.get_flow(flow_id):
        raise HTTPException(404, "自动化流程不存在")
    return store.list_flow_versions(flow_id)


@app.post("/api/flows/{flow_id}/checkpoints")
def create_flow_checkpoint(flow_id: str, payload: FlowCheckpointIn) -> dict[str, Any]:
    try:
        return store.create_flow_checkpoint(flow_id, payload.source, payload.label)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc


@app.post("/api/flow-versions/{version_id}/restore")
def restore_flow_version(version_id: str) -> dict[str, Any]:
    value = store.restore_flow_version(version_id)
    if not value:
        raise HTTPException(404, "流程版本不存在")
    return value


@app.delete("/api/flows/{flow_id}")
def delete_flow(flow_id: str) -> dict[str, bool]:
    try:
        if not store.delete_flow(flow_id):
            raise HTTPException(404, "自动化流程不存在")
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    return {"deleted": True}


@app.get("/api/projects/{project_id}/test-cases")
def test_cases(project_id: str, q: str = "", module: str = "", priority: str = "", status: str = "") -> list[dict[str, Any]]:
    return store.list_test_cases(project_id, q, module, priority, status)


@app.get("/api/test-cases/{case_id}")
def test_case(case_id: str) -> dict[str, Any]:
    value = store.get_test_case(case_id)
    if not value:
        raise HTTPException(404, "测试用例不存在")
    return value


@app.get("/api/test-cases/{case_id}/history")
def test_case_history(case_id: str) -> dict[str, Any]:
    value = store.get_test_case(case_id)
    if not value:
        raise HTTPException(404, "测试用例不存在")
    return store.get_test_case_history(case_id)


@app.post("/api/test-cases")
def save_test_case(payload: TestCaseIn) -> dict[str, Any]:
    try:
        return store.save_test_case(payload.model_dump())
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.delete("/api/test-cases/{case_id}")
def delete_test_case(case_id: str) -> dict[str, bool]:
    if not store.delete_test_case(case_id):
        raise HTTPException(404, "测试用例不存在")
    return {"deleted": True}


@app.post("/api/test-cases/batch-delete")
def batch_delete_test_cases(payload: TestCaseBatchDeleteIn) -> dict[str, int]:
    deleted = store.delete_test_cases(payload.project_id, payload.case_ids)
    return {"deleted": deleted, "requested": len(payload.case_ids)}


_IMPORT_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "case_id": (
        "用例编号", "测试用例编号", "测试编号", "案例编号", "用例序号", "用例号", "序号", "编号", "编码", "用例ID", "用例id", "case id", "caseid",
        "case no", "case number", "test case id", "testcaseid", "test case no", "test id", "testid", "id",
    ),
    "title": (
        "用例名称", "用例标题", "测试用例", "测试名称", "测试标题", "测试点", "用例描述", "用例说明", "名称", "标题",
        "用例名", "name", "title", "test case", "testcase", "test name", "description",
    ),
    "module": (
        "测试模块", "所属模块", "功能模块", "业务模块", "模块", "目录", "分类", "子模块", "子目录", "module", "component",
        "feature", "area", "directory", "folder",
    ),
    "priority": (
        "优先级", "优先级别", "级别", "等级", "用例等级", "重要程度", "风险等级", "priority", "priority level",
        "severity", "level",
    ),
    "case_type": ("用例类型", "测试类型", "类型", "case type", "test type", "type"),
    "preconditions": (
        "前置条件", "前提条件", "前置", "前提", "准备条件", "环境前提", "precondition", "preconditions", "prerequisite",
        "prerequisites", "pre", "setup",
    ),
    "steps": (
        "操作步骤", "操作步骤给ai生成流程用", "操作步骤给AI生成流程用", "测试步骤", "执行步骤", "步骤描述", "步骤说明",
        "操作流程", "测试流程", "执行流程", "步骤", "step", "steps", "test step", "test steps", "teststep",
        "teststeps", "operation", "operations", "operation steps", "actions", "action", "action description", "procedure", "scenario",
    ),
    "test_data": (
        "测试数据", "测试参数", "输入数据", "输入参数", "输入值", "参数", "数据", "样例数据", "测试输入", "测试变量", "test data", "testdata",
        "test parameter", "test parameters", "input", "inputs", "input data", "data", "fixture", "fixtures", "variables",
        "账号", "用户名", "密码",
    ),
    "expected": (
        "预期结果", "期望结果", "预期", "期望", "验收标准", "预期表现", "结果", "expected", "expected result", "expected results",
        "expected outcome", "acceptance criteria", "assertion", "expects", "expect",
    ),
    "tags": ("标签", "标记", "分类标签", "tags", "tag", "labels", "label"),
    "requirement_basis": ("需求依据", "需求来源", "需求编号", "关联需求", "requirement", "requirement id", "requirement basis"),
    "status": ("用例状态", "测试状态", "状态", "case status", "test status", "status"),
    "automation_flow_id": ("自动化流程", "流程ID", "流程id", "automation flow", "automation flow id", "flow id", "flowid"),
}


def _normalize_import_header(value: Any) -> str:
    """Normalize Chinese/English headers for tolerant matching."""
    text = unicodedata.normalize("NFKC", str(value or "")).replace("\ufeff", "").strip().lower()
    return re.sub(r"[\s\u3000_\-./\\:：()（）\[\]【】{}<>《》'\"`#]+", "", text)


_IMPORT_ALIAS_LOOKUP = {
    _normalize_import_header(alias): field
    for field, aliases in _IMPORT_FIELD_ALIASES.items()
    for alias in aliases
}


def _canonical_import_field(value: Any, *, fuzzy: bool = True) -> str | None:
    normalized = _normalize_import_header(value)
    if not normalized:
        return None
    direct = _IMPORT_ALIAS_LOOKUP.get(normalized)
    if direct:
        return direct
    # Allow numbered columns such as "步骤1", "Step 2", "预期结果-3".
    if re.search(r"(?:步骤|step|action|操作|执行)(?:[一二三四五六七八九十0-9]+)?$", normalized, re.I):
        return "steps"
    if re.search(r"(?:预期|期望|expected|expect)(?:结果|result|outcome)?(?:[一二三四五六七八九十0-9]+)?$", normalized, re.I):
        return "expected"
    if re.search(r"(?:数据|参数|data|input|fixture)(?:[一二三四五六七八九十0-9]+)?$", normalized, re.I):
        return "test_data"
    if not fuzzy:
        return None
    # Prefixes are more reliable than generic aliases such as "description"
    # (otherwise a header named "Step Description" would be mistaken for a
    # test title).
    if normalized.startswith(("step", "teststep", "操作步骤", "测试步骤", "执行步骤", "步骤", "操作流程", "测试流程", "action", "operation", "procedure", "scenario")):
        return "steps"
    if normalized.startswith(("expected", "expect", "预期", "期望")):
        return "expected"
    if normalized.startswith(("precondition", "prerequisite", "前置", "前提")):
        return "preconditions"
    if normalized.startswith(("testdata", "inputdata", "input", "fixture", "数据", "参数")):
        return "test_data"
    # Long aliases win to avoid a generic "data"/"id" match swallowing another field.
    for alias, field in sorted(_IMPORT_ALIAS_LOOKUP.items(), key=lambda item: len(item[0]), reverse=True):
        # Do not classify arbitrary cell content (for example a step sentence
        # containing "输入") as a header. Fuzzy matching is limited to short,
        # header-shaped strings; exact aliases above remain unrestricted.
        if len(alias) >= 4 and len(normalized) <= max(40, len(alias) + 16) and alias in normalized:
            return field
    return None


def _decode_import_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "gb18030", "cp936"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" not in text[:256]:
            return text
    return content.decode("utf-8", errors="replace")


def _stringify_import_value(value: Any, *, numbered: bool = False) -> str:
    """Convert Excel/JSON arrays and objects to readable persisted text."""
    if value is None:
        return ""
    if isinstance(value, Mapping):
        parts = []
        for key, item in value.items():
            rendered = _stringify_import_value(item)
            if rendered:
                parts.append(f"{key}: {rendered}")
        return "\n".join(parts)
    if isinstance(value, (list, tuple, set)):
        parts = [_stringify_import_value(item) for item in value]
        parts = [item for item in parts if item]
        if numbered:
            return "\n".join(f"{index}. {item.removeprefix(f'{index}. ').strip()}" for index, item in enumerate(parts, 1))
        return "\n".join(parts)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    # CSV/Excel exports sometimes serialize array fields as JSON text. Decode
    # those values so steps and expected results remain readable in the UI.
    if text[:1] in {"[", "{"} and text[-1:] in {"]", "}"}:
        try:
            parsed = json.loads(text)
        except (TypeError, ValueError, json.JSONDecodeError):
            parsed = None
        if isinstance(parsed, (list, tuple, Mapping)):
            return _stringify_import_value(parsed, numbered=numbered)
    return text


def _normalize_import_row(raw: Mapping[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    # Some JSON exports wrap the actual case under fields/case/test_case/attributes.
    expanded: dict[str, Any] = {}
    for key in ("fields", "case", "test_case", "testCase", "attributes"):
        nested = raw.get(key)
        if isinstance(nested, Mapping):
            expanded.update(nested)
    expanded.update(raw)
    for key, value in expanded.items():
        field = _canonical_import_field(key)
        if not field:
            continue
        if field in {"steps", "preconditions", "test_data", "expected"}:
            rendered = _stringify_import_value(value, numbered=field in {"steps", "preconditions", "expected"})
        elif field == "tags":
            rendered = value if isinstance(value, list) else _stringify_import_value(value)
        else:
            rendered = _stringify_import_value(value)
        if not rendered and field in row:
            continue
        if field in {"steps", "preconditions", "test_data", "expected"} and row.get(field):
            row[field] = f"{row[field]}\n{rendered}" if rendered else row[field]
        else:
            row[field] = rendered
    # Common compact JSON exports use pre/expects/level and module+sub.
    if not row.get("preconditions") and raw.get("pre") is not None:
        row["preconditions"] = _stringify_import_value(raw.get("pre"), numbered=True)
    if not row.get("expected") and raw.get("expects") is not None:
        row["expected"] = _stringify_import_value(raw.get("expects"), numbered=True)
    if not row.get("priority") and raw.get("level") is not None:
        row["priority"] = _stringify_import_value(raw.get("level"))
    if row.get("module") and raw.get("sub"):
        row["module"] = f"{row['module']} / {_stringify_import_value(raw.get('sub'))}"
    return row


def _json_import_rows(data: Any) -> list[dict[str, Any]]:
    if isinstance(data, list):
        return [dict(item) for item in data if isinstance(item, Mapping)]
    if not isinstance(data, Mapping):
        raise ValueError("JSON 根节点必须是对象或数组")
    for key in ("rows", "data", "items", "cases", "test_cases", "testCases", "records", "results"):
        value = data.get(key)
        if isinstance(value, list):
            return [dict(item) for item in value if isinstance(item, Mapping)]
        if isinstance(value, Mapping):
            try:
                nested = _json_import_rows(value)
            except ValueError:
                nested = []
            if nested:
                return nested
    # Some exporters use an object keyed by case ID instead of an array.
    values = list(data.values())
    if values and all(isinstance(item, Mapping) for item in values):
        candidates = [dict(item) for item in values if any(_canonical_import_field(key, fuzzy=False) for key in item)]
        if candidates:
            return candidates
    # A single case object is also accepted.
    if _canonical_import_field(next(iter(data), "")) or any(key in data for key in ("title", "name", "steps", "test_steps")):
        return [dict(data)]
    raise ValueError("JSON 中未找到用例数组（支持 rows/data/items/cases/test_cases）")


def _header_row_index(values: list[tuple[Any, ...]], *, fuzzy: bool = True) -> int | None:
    best: tuple[int, int] | None = None
    for index, values_row in enumerate(values[:40]):
        cells = [str(item or "").strip() for item in values_row]
        nonempty = sum(bool(cell) for cell in cells)
        strict_hits = sum(_canonical_import_field(cell, fuzzy=False) is not None for cell in cells)
        hits = strict_hits
        if fuzzy:
            # A second pass accepts header variants such as "Test Case Title"
            # while still requiring two independent header-like cells.
            hits = max(strict_hits, sum(_canonical_import_field(cell) is not None for cell in cells))
        minimum_hits = 1 if not fuzzy else 2
        if hits >= minimum_hits and nonempty >= 2:
            score = hits * 10 + min(nonempty, 8)
            if best is None or score > best[1]:
                best = (index, score)
    return best[0] if best else None


def _unique_headers(values_row: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values_row, 1):
        base = _stringify_import_value(value) or f"column_{index}"
        count = used.get(base, 0) + 1
        used[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _parse_import_rows(filename: str, content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        rows = _json_import_rows(json.loads(_decode_import_text(content)))
        return rows, sorted({str(key) for row in rows for key in row})
    if suffix == ".csv":
        text = _decode_import_text(content)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        values = [tuple(row) for row in csv.reader(io.StringIO(text), dialect=dialect)]
        header_index = _header_row_index(values)
        if header_index is None:
            raise ValueError("CSV 中未找到可识别的用例表头")
        headers = _unique_headers(values[header_index])
        rows: list[dict[str, Any]] = []
        for row_values in values[header_index + 1 :]:
            if not any(str(value or "").strip() for value in row_values):
                continue
            if _header_row_index([tuple(row_values)], fuzzy=False) == 0:
                continue
            rows.append(dict(zip(headers, row_values, strict=False)))
        return rows, headers
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("导入 Excel 需要安装 openpyxl") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        all_headers: list[str] = []
        for sheet in workbook.worksheets:
            values = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            if not values:
                continue
            header_index = _header_row_index(values)
            if header_index is None:
                continue
            headers = _unique_headers(values[header_index])
            for row_values in values[header_index + 1 :]:
                if all(value is None or str(value).strip() == "" for value in row_values):
                    continue
                # Repeated header rows inside a sheet are ignored.
                if _header_row_index([tuple(row_values)], fuzzy=False) == 0:
                    continue
                rows.append(dict(zip(headers, row_values, strict=False)))
            for header in headers:
                if header not in all_headers:
                    all_headers.append(header)
        return rows, all_headers
    raise ValueError("仅支持 Excel(.xlsx/.xlsm)、CSV、JSON 文件；旧版 .xls 请另存为 .xlsx")


def _normalize_priority(value: Any) -> str:
    text = _stringify_import_value(value).strip().upper()
    aliases = {
        "0": "P0", "P0": "P0", "紧急": "P0", "阻断": "P0", "CRITICAL": "P0", "BLOCKER": "P0",
        "1": "P1", "P1": "P1", "高": "P1", "HIGH": "P1", "严重": "P1", "MAJOR": "P1",
        "2": "P2", "P2": "P2", "中": "P2", "MEDIUM": "P2", "一般": "P2", "NORMAL": "P2",
        "3": "P3", "P3": "P3", "低": "P3", "LOW": "P3", "次要": "P3", "MINOR": "P3",
    }
    return aliases.get(text, text or "P1")


def _validate_import_rows(rows: list[dict[str, Any]], existing_case_ids: set[str] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[str] = set()
    existing_case_ids = {str(item).strip().upper() for item in (existing_case_ids or set()) if str(item).strip()}
    for index, raw in enumerate(rows, 2):
        row = _normalize_import_row(raw)
        row["case_id"] = _stringify_import_value(row.get("case_id")).strip()
        row["title"] = _stringify_import_value(row.get("title")).strip()
        row["module"] = _stringify_import_value(row.get("module")).strip()
        row["priority"] = _normalize_priority(row.get("priority"))
        row["case_type"] = _stringify_import_value(row.get("case_type")).strip() or "回归测试"
        row["preconditions"] = _stringify_import_value(row.get("preconditions"))
        row["steps"] = _stringify_import_value(row.get("steps"), numbered=True)
        row["test_data"] = _stringify_import_value(row.get("test_data"))
        row["expected"] = _stringify_import_value(row.get("expected"), numbered=True)
        row["status"] = _stringify_import_value(row.get("status")).strip() or "未自动化"
        if isinstance(row.get("tags"), list):
            row["tags"] = [_stringify_import_value(item).strip() for item in row["tags"] if _stringify_import_value(item).strip()]
        else:
            row["tags"] = [item.strip() for item in re.split(r"[,，;；\n]", _stringify_import_value(row.get("tags"))) if item.strip()]
        problems: list[str] = []
        case_key = row["case_id"].upper()
        if not row["title"]:
            problems.append("用例名称为空")
        # Empty IDs are intentionally allowed; the store generates a stable ID on save.
        if case_key and case_key in seen:
            problems.append("用例编号重复")
        if case_key and case_key in existing_case_ids:
            problems.append("用例编号已存在")
        if row["priority"] not in {"P0", "P1", "P2", "P3"}:
            problems.append("优先级非法（支持 P0/P1/P2/P3 或 高/中/低）")
        if problems:
            errors.append({"row": index, "errors": problems, "data": row})
        else:
            if case_key:
                seen.add(case_key)
            valid.append(row)
    return valid, errors


@app.post("/api/test-cases/import/preview")
def preview_test_case_import(payload: ImportPreviewIn) -> dict[str, Any]:
    try:
        rows, headers = _parse_import_rows(payload.filename, base64.b64decode(payload.content_base64))
        strategy = "overwrite" if payload.duplicate_strategy.lower() == "overwrite" else "skip"
        existing = set() if strategy == "overwrite" else {item["case_id"] for item in store.list_test_cases(payload.project_id) if item.get("case_id")}
        valid, errors = _validate_import_rows(rows, existing)
        return {"headers": headers or sorted({key for row in rows for key in row}), "rows": valid, "errors": errors, "total": len(rows), "duplicate_strategy": strategy}
    except (ValueError, json.JSONDecodeError, base64.binascii.Error) as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/test-cases/import/confirm")
def confirm_test_case_import(payload: ImportConfirmIn) -> dict[str, int]:
    imported = 0
    skipped = 0
    existing_project_ids = {str(item["case_id"]).strip().upper() for item in store.list_test_cases(payload.project_id) if item.get("case_id")}
    # Normalize again at the persistence boundary so callers that skip the
    # preview endpoint cannot accidentally write list/dict values into SQLite.
    normalized_rows, _ = _validate_import_rows(payload.rows, set())
    strategy = "overwrite" if payload.duplicate_strategy.lower() == "overwrite" else "skip"
    for row in normalized_rows:
        row["project_id"] = payload.project_id
        try:
            case_key = str(row.get("case_id") or "").strip().upper()
            if case_key and case_key in existing_project_ids:
                if strategy == "overwrite":
                    store.save_test_case(row)
                    imported += 1
                else:
                    skipped += 1
            else:
                saved = store.save_test_case(row)
                imported += 1
                saved_key = str(saved.get("case_id") or row.get("case_id") or "").strip().upper()
                if saved_key:
                    existing_project_ids.add(saved_key)
        except ValueError:
            skipped += 1
    return {"imported": imported, "skipped": skipped + max(0, len(payload.rows) - len(normalized_rows))}


@app.get("/api/projects/{project_id}/regression-plans")
def regression_plans(project_id: str) -> list[dict[str, Any]]:
    return store.list_regression_plans(project_id)


@app.get("/api/regression-plans/{plan_id}")
def regression_plan(plan_id: str) -> dict[str, Any]:
    value = store.get_regression_plan(plan_id)
    if not value: raise HTTPException(404, "回归计划不存在")
    return value


@app.post("/api/regression-plans")
def save_regression_plan(payload: RegressionPlanIn) -> dict[str, Any]:
    try:
        return store.save_regression_plan(payload.model_dump())
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.delete("/api/regression-plans/{plan_id}")
def delete_regression_plan(plan_id: str) -> dict[str, bool]:
    if not store.delete_regression_plan(plan_id): raise HTTPException(404, "回归计划不存在")
    return {"deleted": True}


@app.get("/api/regression-runs")
def regression_runs(plan_id: str | None = None) -> list[dict[str, Any]]:
    return store.list_regression_runs(plan_id)


@app.get("/api/regression-runs/{regression_run_id}")
def regression_run(regression_run_id: str) -> dict[str, Any]:
    value = store.get_regression_run(regression_run_id)
    if not value:
        raise HTTPException(404, "回归执行不存在")
    value["cases"] = store.list_regression_run_cases(regression_run_id)
    return value


@app.get("/api/regression-runs/{regression_run_id}/cases")
def regression_run_cases(regression_run_id: str) -> list[dict[str, Any]]:
    if not store.get_regression_run(regression_run_id):
        raise HTTPException(404, "回归执行不存在")
    return store.list_regression_run_cases(regression_run_id)


@app.get("/api/regression-runs/{regression_run_id}/status")
def regression_run_status(regression_run_id: str) -> dict[str, Any]:
    value = store.get_regression_run(regression_run_id)
    if not value:
        raise HTTPException(404, "回归执行不存在")
    value["cases"] = store.list_regression_run_cases(regression_run_id)
    return value


@app.post("/api/regression-plans/{plan_id}/runs")
def start_regression(plan_id: str) -> dict[str, Any]:
    plan = store.get_regression_plan(plan_id)
    if not plan: raise HTTPException(404, "回归计划不存在")
    if not plan["cases"]: raise HTTPException(400, "回归计划没有可执行的自动化用例")
    environment = next((item for item in store.list_environments(plan["project_id"]) if item["id"] == plan.get("environment_id")), {})
    blocked = [item for item in validate_environment(environment) if item.get("severity") == "block"]
    if blocked:
        raise HTTPException(403, "；".join(item["message"] for item in blocked))
    run = store.create_regression_run(plan)
    active_regressions[run["regression_run_id"]] = threading.Event()
    threading.Thread(target=_run_regression, args=(run, plan), daemon=True).start()
    return run


@app.post("/api/regression-runs/{regression_run_id}/rerun-failed")
def rerun_failed_regression(regression_run_id: str) -> dict[str, Any]:
    previous = store.get_regression_run(regression_run_id)
    if not previous:
        raise HTTPException(404, "回归执行不存在")
    failed = [item for item in store.list_regression_run_cases(regression_run_id) if item.get("status") == "FAIL"]
    if not failed:
        raise HTTPException(400, "上一轮没有失败用例")
    plan = store.get_regression_plan(previous["regression_plan_id"])
    if not plan:
        raise HTTPException(404, "原回归计划不存在")
    plan = dict(plan)
    plan["name"] = f"{plan['name']}（失败重跑）"
    plan["cases"] = [{"case_id": item["case_id"], "title": item.get("title", ""), "priority": item.get("priority", ""), "automation_flow_id": item.get("automation_flow_id"), "enabled": 1} for item in failed]
    run = store.create_regression_run(plan)
    active_regressions[run["regression_run_id"]] = threading.Event()
    threading.Thread(target=_run_regression, args=(run, plan), daemon=True).start()
    return run


@app.post("/api/regression-runs/{regression_run_id}/stop")
def stop_regression(regression_run_id: str) -> dict[str, str]:
    event = active_regressions.get(regression_run_id)
    if not event:
        raise HTTPException(409, "回归执行已结束，无法停止")
    event.set()
    return {"status": "STOPPING"}


@app.get("/api/runs")
def runs(flow_id: str = "", environment_id: str = "", browser: str = "", status: str = "", q: str = "") -> list[dict[str, Any]]:
    values = store.list_runs()
    filtered = [item for item in values if (not flow_id or item.get("flow_id") == flow_id) and (not environment_id or item.get("environment_id") == environment_id) and (not browser or item.get("browser") == browser) and (not status or item.get("status") == status) and (not q or q.lower() in f"{item.get('run_id','')} {item.get('flow_id','')}".lower())]
    return filtered


@app.get("/api/runs/{run_id}/steps")
def run_steps(run_id: str) -> list[dict[str, Any]]:
    if not store.get_run(run_id):
        raise HTTPException(404, "执行记录不存在")
    return store.list_run_steps(run_id)


@app.get("/api/runs/{run_id}/healing-suggestions")
def healing_suggestions(run_id: str, status: str = "") -> list[dict[str, Any]]:
    if not store.get_run(run_id):
        raise HTTPException(404, "执行记录不存在")
    return store.list_heal_suggestions(run_id, status)


@app.post("/api/healing-suggestions/{suggestion_id}/approve")
def approve_healing_suggestion(suggestion_id: str) -> dict[str, Any]:
    value = store.decide_heal_suggestion(suggestion_id, "approved")
    if not value:
        raise HTTPException(404, "自愈建议不存在")
    return value


@app.post("/api/healing-suggestions/{suggestion_id}/reject")
def reject_healing_suggestion(suggestion_id: str) -> dict[str, Any]:
    value = store.decide_heal_suggestion(suggestion_id, "rejected")
    if not value:
        raise HTTPException(404, "自愈建议不存在")
    return value


@app.get("/api/runs/{run_id}/status")
def run_status(run_id: str) -> dict[str, Any]:
    value = store.get_run(run_id)
    if not value:
        raise HTTPException(404, "执行记录不存在")
    executor = active_executors.get(run_id)
    if executor:
        value["status"] = executor.state
    return value


@app.post("/api/runs")
def start_run(payload: RunIn) -> dict[str, Any]:
    flow = store.get_flow(payload.flow_id)
    if not flow:
        raise HTTPException(404, "自动化流程不存在")
    selected_environment_id = payload.environment_id or flow.get("environment_id")
    environment = next((item for item in store.list_environments(flow["project_id"]) if item["id"] == selected_environment_id), {})
    if selected_environment_id and not environment:
        raise HTTPException(400, "测试环境不存在或不属于当前项目")
    blocked = [item for item in validate_environment(environment) if item.get("severity") == "block"]
    if blocked:
        raise HTTPException(403, "；".join(item["message"] for item in blocked))
    run = store.create_run(payload.flow_id, selected_environment_id, payload.browser, payload.dataset_name)
    flow["elements"] = {item["element_id"]: item for item in store.list_elements(flow["project_id"])}
    flow["project_variables"] = store.get_project_variables(flow["project_id"])
    environment["browser"] = payload.browser
    run["from_step"] = payload.from_step
    run["to_step"] = payload.to_step
    run["breakpoints"] = payload.breakpoints or flow.get("breakpoints", [])
    run["runtime_variables"] = payload.runtime_variables
    run["dataset_name"] = payload.dataset_name
    run["ai_config"] = payload.ai
    if run["to_step"] is not None and run["to_step"] < run["from_step"]:
        raise HTTPException(400, "执行范围结束步骤不能小于开始步骤")
    if run["from_step"] > len(flow.get("steps", [])):
        raise HTTPException(400, "开始步骤超出流程范围")
    threading.Thread(target=_run_flow, args=(run, flow, environment), daemon=True).start()
    return run


@app.post("/api/recordings")
def start_recording(payload: RecordingIn) -> dict[str, Any]:
    recording_id = f"recording-{uuid.uuid4().hex[:10]}"
    recorder = PlaywrightRecorder(ROOT / "runs", event=lambda event, value: _publish_threadsafe({"type": event, "recording_id": recording_id, **value}))
    active_recorders[recording_id] = recorder
    try:
        recorder.start(payload.url, recording_id)
    except Exception:
        active_recorders.pop(recording_id, None)
        raise
    return {"recording_id": recording_id, "status": "RECORDING", "url": payload.url}


@app.post("/api/recordings/{recording_id}/stop")
def stop_recording(recording_id: str, project_id: str, flow_id: str | None = None) -> dict[str, Any]:
    recorder = active_recorders.pop(recording_id, None)
    if not recorder:
        raise HTTPException(404, "录制任务不存在或已结束")
    result = recorder.stop()
    saved_elements = []
    for element in result["elements"]:
        element["project_id"] = project_id
        saved_elements.append(store.upsert_element(element))
    if flow_id:
        existing = store.get_flow(flow_id)
        if existing:
            existing["steps"] = result["steps"]
            store.save_flow(existing)
    result["elements"] = saved_elements
    return result


def _run_flow(run: dict[str, Any], flow: dict[str, Any], environment: dict[str, Any]) -> None:
    run_id = run["run_id"]
    executor = PlaywrightExecutor(ROOT / "runs", event=lambda event, value: _executor_event(run_id, event, value), ai_config=_resolve_ai_config(run.get("ai_config")))
    executor.configure(run.get("from_step", 1), run.get("to_step"), run.get("breakpoints", []))
    active_executors[run_id] = executor
    try:
        status, report_path, trace_path = _execute_and_report(run, flow, environment, executor)
        _publish_threadsafe({"type": "run_finished", "run_id": run_id, "status": status, "report_path": str(report_path), "trace_path": trace_path})
    except Exception as exc:  # noqa: BLE001 - service boundary reports all execution errors.
        error = str(exc)
        run_dir = ROOT / "runs" / run_id
        report_path = write_html_report({"run_id": run_id, "flow_name": flow.get("name", "自动化流程"), "browser": run.get("browser", "chromium"), "environment": environment.get("name", ""), "base_url": environment.get("base_url", ""), "started_at": run.get("started_at", ""), "ended_at": now_iso(), "error": error}, [{"step_no": "", "name": "执行器", "action_label": "执行失败", "status": "FAIL", "error_message": error}], run_dir / "report.html")
        store.finish_run(run_id, "FAIL", str(report_path), "", error)
        _publish_threadsafe({"type": "run_finished", "run_id": run_id, "status": "FAIL", "report_path": str(report_path), "error": error})
    finally:
        active_executors.pop(run_id, None)


def _execute_and_report(run: dict[str, Any], flow: dict[str, Any], environment: dict[str, Any], executor: PlaywrightExecutor) -> tuple[str, Path, str]:
    run_id = run["run_id"]
    status, step_results, trace_path = executor.run(flow, environment, run_id, run.get("runtime_variables", {}))
    persisted_ids = {item.get("step_id") for item in store.list_run_steps(run_id)}
    for result in step_results:
        if result.get("step_id") in persisted_ids:
            continue
        _create_healing_suggestion(run_id, result)
        run_step_id = store.add_run_step(run_id, result)
        if result.get("step_id"):
            source_step = next((item for item in flow.get("steps", []) if item.get("id") == result["step_id"]), None)
            if source_step and source_step.get("element_id"):
                store.increment_element_usage(source_step["element_id"])
        if result.get("screenshot_path"):
            store.add_artifact(run_id, "screenshot", result["screenshot_path"], run_step_id)
        if result.get("before_screenshot_path"):
            store.add_artifact(run_id, "before_screenshot", result["before_screenshot_path"], run_step_id)
        if result.get("after_screenshot_path"):
            store.add_artifact(run_id, "after_screenshot", result["after_screenshot_path"], run_step_id)
        if result.get("log_path"):
            store.add_artifact(run_id, "log", result["log_path"], run_step_id)
    report_run = {"run_id": run_id, "flow_name": flow["name"], "browser": run["browser"], "environment": environment.get("name", ""), "base_url": environment.get("base_url", ""), "started_at": run.get("started_at", ""), "ended_at": now_iso(), "trace_path": trace_path}
    report_path = write_html_report(report_run, step_results, ROOT / "runs" / run_id / "report.html")
    store.add_artifact(run_id, "trace", trace_path)
    store.add_artifact(run_id, "report", str(report_path))
    store.finish_run(run_id, status, str(report_path), trace_path)
    return status, report_path, trace_path


def _mission_phase(mission_id: str, phase: str, progress: int, message: str) -> None:
    store.update_mission(mission_id, status=phase, progress=progress, started=phase == "DATA_PREPARATION")
    store.add_mission_event(mission_id, "phase_changed", phase, message, {"progress": progress})
    _publish_threadsafe({"type": "mission_updated", "mission_id": mission_id, "status": phase, "progress": progress, "message": message})


def _run_mission(mission_id: str, stop_event: threading.Event) -> None:
    try:
        mission = store.get_mission(mission_id)
        if not mission:
            return
        environment = next(
            (item for item in store.list_environments(mission["project_id"]) if item["id"] == mission.get("environment_id")),
            None,
        )
        if not environment:
            store.update_mission(mission_id, status="REPORT_COMPLETE", verdict="BLOCKED", progress=100,
                                 blocking_reason="测试环境不存在", ended=True)
            return

        _mission_phase(mission_id, "DATA_PREPARATION", 25, "正在确认数据边界与执行命名空间")
        store.add_mission_event(mission_id, "data_policy", "DATA_PREPARATION",
                                "本轮未创建业务数据；不会修改或删除来源不明的数据",
                                {"allowed_ownership": ["KLEIN_OWNED", "USER_APPROVED"]})
        if stop_event.is_set():
            store.update_mission(mission_id, status="REPORT_COMPLETE", verdict="CANCELLED", progress=100, ended=True)
            return

        _mission_phase(mission_id, "EXECUTING", 35, "开始执行已通过前置检查的场景")
        scenarios = mission.get("scenarios", [])
        verdicts: list[str] = []
        for index, scenario in enumerate(scenarios):
            if stop_event.is_set():
                store.update_mission_scenario(scenario["scenario_id"], status="CANCELLED", verdict="CANCELLED", ended=True)
                verdicts.append("CANCELLED")
                continue
            case = store.get_test_case(str(scenario.get("case_id") or ""))
            flow = store.get_flow(str((case or {}).get("automation_flow_id") or ""))
            if not case or not flow:
                store.update_mission_scenario(scenario["scenario_id"], status="COMPLETED", verdict="BLOCKED",
                                              findings=[{"code": "EXECUTION_ADAPTER_MISSING", "message": "没有可用的执行适配器"}], ended=True)
                verdicts.append("BLOCKED")
                continue

            store.update_mission_scenario(scenario["scenario_id"], status="RUNNING", started=True)
            flow["elements"] = {item["element_id"]: item for item in store.list_elements(flow["project_id"])}
            flow["project_variables"] = store.get_project_variables(flow["project_id"])
            execution_environment = {**environment, "browser": environment.get("browser", "chromium")}
            run = store.create_run(flow["flow_id"], mission.get("environment_id"), execution_environment["browser"], f"Mission {mission_id}")
            store.update_mission_scenario(scenario["scenario_id"], run_id=run["run_id"])
            executor = PlaywrightExecutor(ROOT / "runs", event=lambda event, value, run_id=run["run_id"]: _executor_event(run_id, event, {"mission_id": mission_id, "scenario_id": scenario["scenario_id"], **value}), ai_config=_resolve_ai_config())
            active_executors[run["run_id"]] = executor
            try:
                _execute_and_report(run, flow, execution_environment, executor)
                persisted_run = store.get_run(run["run_id"]) or run
                result = verify_execution(flow.get("steps", []), store.list_run_steps(run["run_id"]), str(persisted_run.get("trace_path") or ""), str(scenario.get("risk") or "P1"))
            except Exception as exc:  # noqa: BLE001
                error = str(exc)
                store.finish_run(run["run_id"], "FAIL", error_message=error)
                result = {
                    "verdict": "BLOCKED",
                    "ledger": [],
                    "verifications": [{"layer": "runtime", "status": "BLOCKED", "rule": "执行器必须正常完成", "observed": error, "evidence": []}],
                }
            finally:
                active_executors.pop(run["run_id"], None)

            store.save_mission_verdict(mission_id, scenario["scenario_id"], result)
            findings = [
                {"layer": item.get("layer"), "status": item.get("status"), "message": item.get("observed", "")}
                for item in result.get("verifications", []) if item.get("status") != "PASS"
            ]
            store.update_mission_scenario(scenario["scenario_id"], status="COMPLETED", verdict=result["verdict"], findings=findings, ended=True)
            verdicts.append(result["verdict"])
            progress = 35 + round(((index + 1) / max(1, len(scenarios))) * 40)
            store.update_mission(mission_id, progress=progress)
            store.add_mission_event(mission_id, "scenario_finished", "EXECUTING",
                                    f"{scenario['title']}：{result['verdict']}", {"scenario_id": scenario["scenario_id"], "verdict": result["verdict"], "run_id": run["run_id"]})
            _publish_threadsafe({"type": "mission_scenario_finished", "mission_id": mission_id, "scenario_id": scenario["scenario_id"], "verdict": result["verdict"], "progress": progress})

        _mission_phase(mission_id, "VERIFYING", 80, "独立验证器正在汇总覆盖账本与证据")
        final_verdict = aggregate_verdict(verdicts, stop_event.is_set())
        if "FAIL" in verdicts and not stop_event.is_set():
            _mission_phase(mission_id, "FAILURE_REPRODUCTION", 86, "正在使用相同环境与流程复现失败项")
            latest_mission = store.get_mission(mission_id) or mission
            failed_scenarios = [item for item in latest_mission.get("scenarios", []) if item.get("verdict") == "FAIL"]
            for failed_scenario in failed_scenarios:
                if stop_event.is_set():
                    break
                case = store.get_test_case(str(failed_scenario.get("case_id") or ""))
                flow = store.get_flow(str((case or {}).get("automation_flow_id") or ""))
                original_steps = store.list_run_steps(str(failed_scenario.get("run_id") or ""))
                if not case or not flow:
                    classification = {"classification": "UNCONFIRMED", "confidence": 40, "message": "复现时找不到来源用例或执行流程"}
                    store.update_mission_scenario(failed_scenario["scenario_id"], failure_classification=classification["classification"], reproduction_confidence=classification["confidence"])
                    continue
                flow["elements"] = {item["element_id"]: item for item in store.list_elements(flow["project_id"])}
                flow["project_variables"] = store.get_project_variables(flow["project_id"])
                reproduction = store.create_run(flow["flow_id"], mission.get("environment_id"), environment.get("browser", "chromium"), f"Mission {mission_id} 失败复现")
                executor = PlaywrightExecutor(ROOT / "runs", event=lambda event, value, run_id=reproduction["run_id"], scenario_id=failed_scenario["scenario_id"]: _executor_event(run_id, event, {"mission_id": mission_id, "scenario_id": scenario_id, "reproduction": True, **value}), ai_config=_resolve_ai_config())
                active_executors[reproduction["run_id"]] = executor
                try:
                    _execute_and_report(reproduction, flow, environment, executor)
                    persisted = store.get_run(reproduction["run_id"]) or reproduction
                    reproduced_steps = store.list_run_steps(reproduction["run_id"])
                    reproduced_result = verify_execution(flow.get("steps", []), reproduced_steps, str(persisted.get("trace_path") or ""), str(failed_scenario.get("risk") or "P1"))
                except Exception as exc:  # noqa: BLE001
                    store.finish_run(reproduction["run_id"], "FAIL", error_message=str(exc))
                    reproduced_steps = store.list_run_steps(reproduction["run_id"])
                    reproduced_result = {"verdict": "BLOCKED"}
                finally:
                    active_executors.pop(reproduction["run_id"], None)
                classification = classify_reproduction(original_steps, reproduced_steps, str(reproduced_result.get("verdict") or "INCONCLUSIVE"))
                tapd_draft = build_tapd_draft(case, reproduced_steps or original_steps) if classification["classification"] == "PRODUCT_DEFECT_CANDIDATE" else {}
                store.update_mission_scenario(
                    failed_scenario["scenario_id"], reproduction_run_id=reproduction["run_id"],
                    reproduction_verdict=str(reproduced_result.get("verdict") or "INCONCLUSIVE"),
                    failure_classification=classification["classification"],
                    reproduction_confidence=int(classification["confidence"]), tapd_draft=tapd_draft,
                )
                store.add_mission_event(mission_id, "failure_reproduced", "FAILURE_REPRODUCTION", classification["message"], {
                    "scenario_id": failed_scenario["scenario_id"], "original_run_id": failed_scenario.get("run_id", ""),
                    "reproduction_run_id": reproduction["run_id"], "reproduction_verdict": reproduced_result.get("verdict"),
                    "classification": classification["classification"], "confidence": classification["confidence"],
                })
        _mission_phase(mission_id, "DATA_CLEANUP", 92, "正在核对数据所有权与清理结果")
        store.add_mission_event(mission_id, "cleanup_finished", "DATA_CLEANUP",
                                "本轮没有登记 KLEIN_OWNED 数据，无需执行删除", {"deleted": 0, "retained": 0})
        store.update_mission(mission_id, status="REPORT_COMPLETE", verdict=final_verdict, progress=100,
                             blocking_reason="" if final_verdict not in {"BLOCKED", "INCONCLUSIVE"} else "查看场景验证与覆盖账本", ended=True)
        store.add_mission_event(mission_id, "mission_finished", "REPORT_COMPLETE", f"任务完成：{final_verdict}", {"verdict": final_verdict})
        _publish_threadsafe({"type": "mission_finished", "mission_id": mission_id, "status": "REPORT_COMPLETE", "verdict": final_verdict, "progress": 100})
    except Exception as exc:  # noqa: BLE001
        store.update_mission(mission_id, status="REPORT_COMPLETE", verdict="BLOCKED", progress=100,
                             blocking_reason=str(exc), ended=True)
        store.add_mission_event(mission_id, "mission_error", "REPORT_COMPLETE", str(exc))
        _publish_threadsafe({"type": "mission_finished", "mission_id": mission_id, "status": "REPORT_COMPLETE", "verdict": "BLOCKED", "error": str(exc), "progress": 100})
    finally:
        active_missions.pop(mission_id, None)


def _run_regression(regression_run: dict[str, Any], plan: dict[str, Any]) -> None:
    regression_id = regression_run["regression_run_id"]
    stop_event = active_regressions.get(regression_id) or threading.Event()
    environment = next((item for item in store.list_environments(plan["project_id"]) if item["id"] == plan.get("environment_id")), {})
    environment["browser"] = plan.get("browser", "chromium")
    counts = {"PASS": 0, "FAIL": 0, "SKIP": 0, "BLOCKED": 0}
    cases_for_report: list[dict[str, Any]] = []
    for order, case in enumerate(plan.get("cases", [])):
        if stop_event.is_set():
            for rest_order, rest in enumerate(plan.get("cases", [])[order:], order):
                counts["SKIP"] += 1
                store.add_regression_case(regression_id, rest["case_id"], None, rest_order, "SKIP", "用户请求停止回归")
                cases_for_report.append({"case_id": rest["case_id"], "title": rest.get("title", ""), "status": "SKIP", "error_message": "用户请求停止回归"})
            break
        if not case.get("enabled", 1):
            counts["SKIP"] += 1
            store.add_regression_case(regression_id, case["case_id"], None, order, "SKIP", "用例已禁用")
            cases_for_report.append({"case_id": case["case_id"], "title": case.get("title", ""), "status": "SKIP"})
            continue
        flow = store.get_flow(case.get("automation_flow_id"))
        if not flow:
            counts["BLOCKED"] += 1
            store.add_regression_case(regression_id, case["case_id"], None, order, "BLOCKED", "绑定的自动化流程不存在")
            cases_for_report.append({"case_id": case["case_id"], "title": case.get("title", ""), "status": "BLOCKED", "error_message": "绑定的自动化流程不存在"})
            continue
        flow["elements"] = {item["element_id"]: item for item in store.list_elements(flow["project_id"])}
        flow["project_variables"] = store.get_project_variables(flow["project_id"])
        run = store.create_run(flow["flow_id"], plan.get("environment_id"), plan.get("browser", "chromium"))
        executor = PlaywrightExecutor(ROOT / "runs", event=lambda event, value: _executor_event(run["run_id"], event, {"regression_run_id": regression_id, **value}), ai_config=_resolve_ai_config())
        active_executors[run["run_id"]] = executor
        try:
            status, report_path, trace_path = _execute_and_report(run, flow, environment, executor)
            error_message = ""
            if status != "PASS":
                failed_steps = store.list_run_steps(run["run_id"])
                error_message = next((item.get("error_message") or item.get("message") for item in failed_steps if item.get("status") == "FAIL"), "执行失败")
        except Exception as exc:  # noqa: BLE001
            status, report_path, trace_path = "FAIL", Path(""), ""
            error_message = str(exc)
            failure_report = write_html_report({"run_id": run["run_id"], "flow_name": flow.get("name", "自动化流程"), "browser": run.get("browser", "chromium"), "environment": environment.get("name", ""), "base_url": environment.get("base_url", ""), "started_at": run.get("started_at", ""), "ended_at": now_iso()}, [{"step_no": "", "name": "执行器", "action_label": "执行失败", "status": "FAIL", "error_message": error_message}], ROOT / "runs" / run["run_id"] / "report.html")
            report_path = failure_report
            store.finish_run(run["run_id"], status, str(failure_report), "")
        finally:
            active_executors.pop(run["run_id"], None)
        counts[status] = counts.get(status, 0) + 1
        store.add_regression_case(regression_id, case["case_id"], run["run_id"], order, status, error_message)
        cases_for_report.append({"case_id": case["case_id"], "title": case.get("title", ""), "status": status, "run_id": run["run_id"], "report_path": str(report_path), "trace_path": trace_path, "error_message": error_message})
        _publish_threadsafe({"type": "regression_case_finished", "regression_run_id": regression_id, "case_id": case["case_id"], "run_id": run["run_id"], "status": status})
        if status != "PASS" and plan.get("stop_on_failure"):
            for rest_order, rest in enumerate(plan.get("cases", [])[order + 1:], order + 1):
                if rest.get("enabled", 1):
                    counts["SKIP"] += 1
                    store.add_regression_case(regression_id, rest["case_id"], None, rest_order, "SKIP", "前置用例失败，计划设置为失败即停止")
                    cases_for_report.append({"case_id": rest["case_id"], "title": rest.get("title", ""), "status": "SKIP", "error_message": "前置用例失败，计划设置为失败即停止"})
            break
    overall = "STOPPED" if stop_event.is_set() else ("FAIL" if counts["FAIL"] or counts["BLOCKED"] else "PASS")
    report_path = write_regression_report({"regression_run_id": regression_id, "plan_name": plan["name"], "environment": environment.get("name", ""), "browser": plan.get("browser", "chromium"), "started_at": regression_run.get("started_at", ""), "ended_at": now_iso()}, cases_for_report, ROOT / "runs" / regression_id / "report.html")
    store.finish_regression_run(regression_id, overall, counts, str(report_path))
    _publish_threadsafe({"type": "regression_finished", "regression_run_id": regression_id, "status": overall, "report_path": str(report_path), "counts": counts})
    active_regressions.pop(regression_id, None)


def _auto_generate_batch(project_id: str, payload: AiBatchIn) -> None:
    from .ai_client import build_flow_payload, generate_flow_steps
    cases = store.list_test_cases(project_id)
    to_process = [case for case in cases if not (payload.only_unbound and case.get("automation_flow_id"))]
    total = len(to_process)
    state = {"done": 0, "total": total, "failed": 0, "current": "", "finished": False, "failures": []}
    auto_generation_state[project_id] = state
    environments = store.list_environments(project_id)
    environment = environments[0] if environments else {}
    start_url = environment.get("base_url", "")
    for index, case in enumerate(to_process, 1):
        title = case.get("title", "") or case.get("case_id", "")
        state["current"] = title
        try:
            steps = generate_flow_steps(payload.base_url, payload.api_key, payload.model, case, start_url)
            flow_payload = build_flow_payload(case, steps, start_url, environment.get("id"))
            for element in flow_payload.pop("elements", []):
                store.upsert_element(element)
            flow = store.save_flow(flow_payload)
            store.save_test_case({**case, "automation_flow_id": flow["flow_id"], "status": "已自动化"})
            status, error = "ok", ""
        except Exception as exc:  # noqa: BLE001 - batch continues on per-case failure and reports it.
            status, error = "fail", str(exc)
            state["failed"] += 1
            state["failures"].append({"case_id": case.get("case_id", ""), "title": title, "error": error})
        state["done"] = index
        _publish_threadsafe({"type": "auto_generate_progress", "project_id": project_id, "done": index,
                             "total": total, "failed": state["failed"], "case_id": case.get("case_id", ""),
                             "title": title, "status": status, "error": error})
    state["finished"] = True
    state["current"] = ""
    active_auto_generations.pop(project_id, None)
    _publish_threadsafe({"type": "auto_generate_finished", "project_id": project_id, "done": total,
                         "total": total, "failed": state["failed"]})


def _publish_threadsafe(payload: dict[str, Any]) -> None:
    if main_loop and main_loop.is_running():
        asyncio.run_coroutine_threadsafe(_publish(payload), main_loop)


def _run_agent(case_id: str, payload: AiGenerateFlowIn) -> None:
    from .agent import run_case_agent
    case = store.get_test_case(case_id)
    if not case:
        return
    environments = store.list_environments(case["project_id"])
    environment = next((item for item in environments if item["id"] == payload.environment_id), environments[0] if environments else {})
    base_url = payload.start_url or environment.get("base_url", "")
    variables = environment.get("variables") or {}
    try:
        safety_environment = {**environment, "base_url": base_url}
        blocked = [item for item in validate_environment(safety_environment) if item.get("severity") == "block"]
        if blocked:
            raise RuntimeError("；".join(item["message"] for item in blocked))
        result = run_case_agent(
            {"base_url": payload.base_url, "api_key": payload.api_key, "model": payload.model},
            case, base_url, variables, headed=True,
            on_step=lambda step: _publish_threadsafe({"type": "agent_step", "case_id": case_id, **step}),
            environment=safety_environment,
        )
        _publish_threadsafe({"type": "agent_finished", "case_id": case_id, "status": result["status"], "reason": result["reason"]})
    except Exception as exc:  # noqa: BLE001
        _publish_threadsafe({"type": "agent_finished", "case_id": case_id, "status": "FAIL", "reason": str(exc)})
    finally:
        active_agents.pop(case_id, None)


def _executor_event(run_id: str, event: str, value: dict[str, Any]) -> None:
    """Persist control state as it changes while also broadcasting it to the UI."""
    if event == "run_state" and value.get("status"):
        store.update_run_status(run_id, str(value["status"]), str(value.get("error", "")))
    if event == "step_finished" and value.get("step_id"):
        existing = {item.get("step_id") for item in store.list_run_steps(run_id)}
        if value.get("step_id") not in existing:
            run_step_id = store.add_run_step(run_id, value)
            _create_healing_suggestion(run_id, value)
            if value.get("screenshot_path"):
                store.add_artifact(run_id, "screenshot", value["screenshot_path"], run_step_id)
            if value.get("before_screenshot_path"):
                store.add_artifact(run_id, "before_screenshot", value["before_screenshot_path"], run_step_id)
            if value.get("after_screenshot_path"):
                store.add_artifact(run_id, "after_screenshot", value["after_screenshot_path"], run_step_id)
            if value.get("log_path"):
                store.add_artifact(run_id, "log", value["log_path"], run_step_id)
    _publish_threadsafe({"type": event, "run_id": run_id, **value})


def _resolve_ai_config(explicit: dict[str, Any] | None = None) -> dict[str, str]:
    """解析 AI 自愈配置：显式传入优先，其次环境变量，都没有则返回空配置（等于关闭自愈）。"""
    from .ai_locator_heal import normalize_ai_config

    config = normalize_ai_config(explicit)
    if config.get("base_url") and config.get("model"):
        return config
    env_config = {
        "base_url": os.environ.get("LODESTAR_AI_BASE_URL", ""),
        "api_key": os.environ.get("LODESTAR_AI_KEY", "") or os.environ.get("LODESTAR_AI_API_KEY", ""),
        "model": os.environ.get("LODESTAR_AI_MODEL", ""),
    }
    return {key: config.get(key) or env_config.get(key) or "" for key in ("base_url", "api_key", "model")}


def _create_healing_suggestion(run_id: str, step: dict[str, Any]) -> None:
    action_spec = step.get("action_spec") or {}
    action = str(action_spec.get("action") or step.get("action") or "")
    element_id = str(action_spec.get("element_id") or "")
    interactive_actions = {"click", "fill", "clear", "select", "select_value", "upload", "hover", "press"}
    locator = step.get("locator") or {}
    if action not in interactive_actions or not element_id:
        return
    # AI 自愈成功（模型在页面上重新定位并已通过命中校验）也要留一条待审批建议
    if locator.get("ai_healed"):
        store.create_heal_suggestion(element_id, run_id, int(step.get("step_no") or 0), locator, str(step.get("screenshot_path") or ""))
        return
    if not locator.get("used_fallback"):
        return
    store.create_heal_suggestion(element_id, run_id, int(step.get("step_no") or 0), locator, str(step.get("screenshot_path") or ""))


async def _publish(payload: dict[str, Any]) -> None:
    stale: list[WebSocket] = []
    for client in clients:
        try:
            await client.send_text(json.dumps(payload, ensure_ascii=False))
        except Exception:
            stale.append(client)
    for client in stale:
        clients.discard(client)


@app.post("/api/runs/{run_id}/pause")
def pause_run(run_id: str) -> dict[str, str]:
    executor = active_executors.get(run_id)
    if not executor:
        raise HTTPException(409, "执行已结束，无法暂停")
    executor.pause()
    return {"status": executor.state}


@app.post("/api/runs/{run_id}/resume")
def resume_run(run_id: str) -> dict[str, str]:
    executor = active_executors.get(run_id)
    if not executor:
        raise HTTPException(409, "执行已结束，无法继续")
    executor.resume()
    return {"status": executor.state}


@app.post("/api/runs/{run_id}/step")
def step_run(run_id: str) -> dict[str, str]:
    executor = active_executors.get(run_id)
    if not executor:
        raise HTTPException(409, "执行已结束，无法单步")
    executor.step_once()
    return {"status": executor.state}


@app.post("/api/runs/{run_id}/breakpoints")
def set_run_breakpoints(run_id: str, payload: dict[str, list[int]]) -> dict[str, Any]:
    executor = active_executors.get(run_id)
    if not executor:
        raise HTTPException(409, "执行已结束，无法设置断点")
    points = payload.get("breakpoints", [])
    executor.set_breakpoints(points)
    return {"status": executor.state, "breakpoints": sorted(executor._breakpoints)}


@app.post("/api/runs/{run_id}/stop")
def stop_run(run_id: str) -> dict[str, str]:
    executor = active_executors.get(run_id)
    if not executor:
        raise HTTPException(409, "执行已结束，无法停止")
    executor.stop()
    return {"status": executor.state}


def _cleanup_cutoff(days: int | None = None, before: str | None = None) -> float:
    if before:
        try:
            return datetime.fromisoformat(before.replace("Z", "+00:00")).timestamp()
        except ValueError as exc:
            raise HTTPException(400, "自定义日期格式应为 YYYY-MM-DD") from exc
    if days not in {7, 30, 90}:
        raise HTTPException(400, "清理范围只支持 7、30、90 天或自定义日期")
    return datetime.now().timestamp() - int(days) * 86400


def _cleanup_inventory(days: int | None = None, before: str | None = None) -> dict[str, Any]:
    cutoff = _cleanup_cutoff(days, before)
    inventory = {"runs": 0, "screenshots": 0, "logs": 0, "traces": 0, "reports": 0, "other_files": 0, "total_files": 0, "skipped_running": 0, "run_ids": []}
    for path in (ROOT / "runs").glob("*"):
        if not path.is_dir() or path.stat().st_mtime >= cutoff:
            continue
        if path.name in active_executors:
            inventory["skipped_running"] += 1
            continue
        inventory["runs"] += 1
        inventory["run_ids"].append(path.name)
        for child in path.rglob("*"):
            if not child.is_file():
                continue
            inventory["total_files"] += 1
            relative = child.relative_to(path).as_posix()
            if relative.startswith("screenshots/"):
                inventory["screenshots"] += 1
            elif relative.startswith("logs/"):
                inventory["logs"] += 1
            elif relative.startswith("trace/"):
                inventory["traces"] += 1
            elif child.name == "report.html":
                inventory["reports"] += 1
            else:
                inventory["other_files"] += 1
    return inventory


@app.post("/api/artifacts/cleanup/preview")
def cleanup_preview(payload: CleanupPreviewIn) -> dict[str, Any]:
    return _cleanup_inventory(payload.days, payload.before)


@app.delete("/api/artifacts/cleanup")
def cleanup_artifacts(days: int | None = None, before: str | None = None) -> dict[str, int]:
    inventory = _cleanup_inventory(days, before)
    removed = 0
    skipped_running = 0
    for path in (ROOT / "runs").glob("*"):
        if path.is_dir() and path.name in active_executors:
            skipped_running += 1
            continue
        if path.is_dir() and path.name in inventory["run_ids"]:
            for child in sorted(path.rglob("*"), reverse=True):
                if child.is_file():
                    child.unlink(missing_ok=True)
                    removed += 1
            for child in sorted(path.rglob("*"), reverse=True):
                if child.is_dir():
                    child.rmdir()
            path.rmdir()
    return {"removed_files": removed, "skipped_running": max(skipped_running, int(inventory["skipped_running"])), "days": days or 0}


@app.websocket("/ws")
async def websocket(websocket: WebSocket) -> None:
    await websocket.accept()
    clients.add(websocket)
    try:
        while True:
            await websocket.receive_text()
    except WebSocketDisconnect:
        clients.discard(websocket)


@app.get("/api/runtime")
def runtime_info() -> dict[str, Any]:
    return CONFIG.public()


# API and artifact routes are declared above this mount, so they keep priority.
# The browser-facing UI is only served in development; the packaged desktop app
# embeds its own frontend assets, so no web frontend is exposed in production.
if not CONFIG.production:
    app.mount("/", StaticFiles(directory=CONFIG.frontend_dist, html=True, check_dir=False), name="frontend")


@app.on_event("startup")
async def capture_main_loop() -> None:
    global main_loop
    main_loop = asyncio.get_running_loop()


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("backend.app:app", host=CONFIG.backend_host, port=CONFIG.backend_port, reload=False)
